import asyncio
import hmac
import importlib.util
import io
import json
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import tempfile
import threading
import time
import unicodedata
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.exceptions import PackageNotFoundError
from docx.shared import Inches, RGBColor
from flask import Flask, abort, flash, g, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, KeepInFrame, Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

import job_import


MODEL_NAME = os.environ.get("OLLAMA_MODEL", "qwen2.5:3b")
OLLAMA_TIMEOUT = int(os.environ.get("OLLAMA_TIMEOUT", "900"))
COPILOT_TIMEOUT = int(os.environ.get("COPILOT_TIMEOUT", "300"))
MAX_REFERENCE_CONTEXT_CHARS = 12000
MAX_STYLE_EXAMPLES = 2
MAX_STYLE_EXAMPLE_CHARS = 1800
OLLAMA_CONTEXT = 4096
OLLAMA_CONTEXT_WITH_EXAMPLES = 6144
MAX_LINKS_PER_IMPORT = 20
MATCH_DOMAIN_HINT = (
    "data, statistiques, mathématiques appliquées, data science, data analyse, data engineering, "
    "modélisation, actuariat, quantitatif, recherche opérationnelle, machine learning, BI"
)
STATUSES = ["A analyser", "A preparer", "Candidature envoyee", "Entretien", "Refus", "Offre recue"]

app = Flask(__name__, instance_relative_config=True)
app.config.update(
    DATABASE=os.environ.get("CANDIDATURES_DATABASE", str(Path(app.instance_path) / "candidatures.sqlite3")),
    SECRET_KEY=os.environ.get("FLASK_SECRET_KEY", secrets.token_hex(32)),
    MAX_CONTENT_LENGTH=10 * 1024 * 1024,
)


def get_db():
    if "db" not in g:
        database_path = Path(app.config["DATABASE"])
        database_path.parent.mkdir(parents=True, exist_ok=True)
        g.db = sqlite3.connect(database_path)
        g.db.row_factory = sqlite3.Row
    return g.db


def close_db(_error=None):
    database = g.pop("db", None)
    if database is not None:
        database.close()


PROFILE_EXTRA_FIELDS = ("languages", "hobbies", "certifications", "achievements")
PROFILE_URL_FIELDS = ("linkedin_url", "portfolio_url")


def normalize_profile_url(value):
    value = (value or "").strip()
    if not value:
        return ""
    if "://" not in value:
        value = "https://" + value
    parsed = urlsplit(value)
    try:
        parsed.port
    except ValueError as error:
        raise ValueError("Le port du lien est invalide.") from error
    if (
        parsed.scheme not in {"http", "https"}
        or not parsed.hostname
        or parsed.username
        or parsed.password
        or any(character.isspace() or ord(character) < 32 for character in value)
    ):
        raise ValueError("Saisis une adresse web valide, par exemple https://linkedin.com/in/ton-profil.")
    return value.rstrip("/")


def profile_list(profile, key):
    """Lignes saisies dans le profil (une par ligne ou separees par des points-virgules)."""
    value = profile[key] if key in profile.keys() else ""
    return [
        item.strip(" -*•\t")
        for line in (value or "").splitlines()
        for item in line.split(";")
        if item.strip(" -*•\t")
    ]


def init_db():
    database = get_db()
    database.executescript(
        """
        CREATE TABLE IF NOT EXISTS profile (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            full_name TEXT NOT NULL DEFAULT '', email TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL DEFAULT '', location TEXT NOT NULL DEFAULT '',
            postal_address TEXT NOT NULL DEFAULT '',
            linkedin_url TEXT NOT NULL DEFAULT '', portfolio_url TEXT NOT NULL DEFAULT '',
            languages TEXT NOT NULL DEFAULT '', hobbies TEXT NOT NULL DEFAULT '',
            certifications TEXT NOT NULL DEFAULT '', achievements TEXT NOT NULL DEFAULT '',
            resume_source TEXT NOT NULL DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS applications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            company TEXT NOT NULL, position TEXT NOT NULL,
            location TEXT NOT NULL DEFAULT '', url TEXT NOT NULL DEFAULT '',
            company_address TEXT NOT NULL DEFAULT '',
            job_description TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'A analyser',
            notes TEXT NOT NULL DEFAULT '', reply TEXT NOT NULL DEFAULT '',
            cv_draft TEXT NOT NULL DEFAULT '', letter_draft TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS reference_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS style_examples (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL CHECK (kind IN ('cv', 'letter')),
            label TEXT NOT NULL DEFAULT '',
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        INSERT OR IGNORE INTO profile (id) VALUES (1);
        """
    )
    profile_columns = {
        row["name"] for row in database.execute("PRAGMA table_info(profile)").fetchall()
    }
    for column in ("postal_address", *PROFILE_EXTRA_FIELDS, *PROFILE_URL_FIELDS):
        if column not in profile_columns:
            database.execute(f"ALTER TABLE profile ADD COLUMN {column} TEXT NOT NULL DEFAULT ''")
    application_columns = {
        row["name"] for row in database.execute("PRAGMA table_info(applications)").fetchall()
    }
    if "company_address" not in application_columns:
        database.execute("ALTER TABLE applications ADD COLUMN company_address TEXT NOT NULL DEFAULT ''")
    if "cv_layout" not in application_columns:
        database.execute("ALTER TABLE applications ADD COLUMN cv_layout TEXT NOT NULL DEFAULT ''")
    if "match_score" not in application_columns:
        database.execute("ALTER TABLE applications ADD COLUMN match_score INTEGER")
    if "match_reason" not in application_columns:
        database.execute("ALTER TABLE applications ADD COLUMN match_reason TEXT NOT NULL DEFAULT ''")
    database.commit()


def get_profile():
    return get_db().execute("SELECT * FROM profile WHERE id = 1").fetchone()


def get_application(application_id):
    return get_db().execute("SELECT * FROM applications WHERE id = ?", (application_id,)).fetchone()


def get_reference_documents():
    return get_db().execute(
        "SELECT id, filename, content, created_at FROM reference_documents ORDER BY id DESC"
    ).fetchall()


def get_style_examples():
    return get_db().execute(
        "SELECT id, kind, label, content, created_at FROM style_examples ORDER BY id DESC"
    ).fetchall()


def select_style_examples(kind, examples):
    selected = [item for item in examples if item["kind"] == kind][:MAX_STYLE_EXAMPLES]
    return [
        {"label": item["label"], "content": item["content"][:MAX_STYLE_EXAMPLE_CHARS].rsplit("\n", 1)[0]
         if len(item["content"]) > MAX_STYLE_EXAMPLE_CHARS else item["content"]}
        for item in selected
    ]


def build_examples_context(kind, examples):
    chosen = select_style_examples(kind, examples)
    if not chosen:
        return ""
    label = "CV" if kind == "cv" else "lettres de motivation"
    blocks = [
        f"Exemple {number} (déjà validé par le candidat pour {item['label'] or 'une autre offre'}) :\n{item['content']}"
        for number, item in enumerate(chosen, start=1)
    ]
    return (
        f"\n\nEXEMPLES DE {label.upper()} VALIDÉS (modèles de qualité)\n"
        "Imite uniquement leur structure, leur ton, leur longueur et leur façon de formuler. "
        "Ils concernent d'autres offres : n'en reprends aucune entreprise cible, aucune phrase telle quelle "
        "et aucun fait, chiffre ou outil absent des sources de référence du candidat.\n\n"
        + "\n\n".join(blocks)
    )


def extract_latex_text(source):
    source = re.sub(r"(?m)^\s*\\(?:documentclass|usepackage|setmainfont|definecolor|pagestyle|setlength|newenvironment|newcommand)\b.*$", "", source)
    source = re.sub(r"(?m)^\s*\\(?:begin|end)\{document\}\s*$", "", source)
    source = re.sub(r"\\color\{[^{}]*\}", "", source)
    source = re.sub(r"\\textcolor\{[^{}]*\}\{([^{}]*)\}", r"\1", source)
    source = re.sub(r"\\cvsection\{([^{}]*)\}", r"\n\1\n", source)
    source = re.sub(
        r"\\cventry\{([^{}]*)\}\{([^{}]*)\}\{([^{}]*)\}",
        r"\n\1\n\2\n\3\n",
        source,
    )
    source = re.sub(r"\\cvskills\{([^{}]*)\}", r"\nCompétences : \1\n", source)
    source = re.sub(r"\\href\{([^{}]*)\}\{([^{}]*)\}", r"\2 (\1)", source)
    source = re.sub(r"\\(?:textbf|textit|emph|underline)\{([^{}]*)\}", r"\1", source)
    source = re.sub(r"\\begin\{[^{}]*\}|\\end\{[^{}]*\}", "", source)
    source = re.sub(r"\\item\b", "\n- ", source)
    source = re.sub(r"\\fontsize\{[^{}]*\}\{[^{}]*\}\\selectfont", "", source)
    source = re.sub(r"\\(?:par|hfill|selectfont)\b", "\n", source)
    source = source.replace(r"\;", " ").replace(r"\|", "|")
    source = re.sub(r"\\[a-zA-Z@]+\*?(?:\[[^\]]*\])?", "", source)
    for escaped, character in (
        (r"\&", "&"), (r"\%", "%"), (r"\$", "$"), (r"\#", "#"),
        (r"\_", "_"), (r"\{", "{"), (r"\}", "}"),
    ):
        source = source.replace(escaped, character)
    source = source.replace("{", "").replace("}", "")
    return "\n".join(line.strip() for line in source.splitlines() if line.strip())


def extract_reference_document(filename, data):
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("Le fichier ne semble pas etre un PDF valide.")
        try:
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        except (PdfReadError, OSError, ValueError) as error:
            raise ValueError("Le PDF n'a pas pu etre lu.") from error
    elif suffix == ".docx":
        try:
            document = Document(io.BytesIO(data))
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as error:
            raise ValueError("Le fichier Word n'a pas pu etre lu.") from error
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables for row in table.rows
        ]
        text = "\n".join(paragraphs + table_text)
    elif suffix == ".tex":
        try:
            text = extract_latex_text(data.decode("utf-8-sig"))
        except UnicodeDecodeError as error:
            raise ValueError("Le fichier LaTeX doit etre en UTF-8.") from error
    elif suffix == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("Le fichier texte doit etre en UTF-8.") from error
    else:
        raise ValueError("Formats acceptes : PDF, DOCX, TEX et TXT.")
    text = text.strip()
    if not text:
        raise ValueError("Aucun texte exploitable n'a ete trouve dans ce document.")
    return text


def build_reference_context(profile, reference_documents):
    sources = []
    primary = (profile["resume_source"] or "").strip()
    if primary:
        sources.append(("CV principal", primary))
    sources.extend(
        (document["filename"], document["content"])
        for document in reference_documents
        if document["content"].strip()
    )
    if not sources:
        return ""

    primary_budget = min(4500, MAX_REFERENCE_CONTEXT_CHARS)
    remaining = MAX_REFERENCE_CONTEXT_CHARS - (primary_budget if primary else 0)
    secondary_count = len(sources) - (1 if primary else 0)
    secondary_budget = remaining // secondary_count if secondary_count else 0
    seen_lines = set()
    blocks = []
    truncated = False
    for index, (label, source) in enumerate(sources):
        budget = primary_budget if index == 0 and primary else secondary_budget
        unique_lines = []
        for line in source.splitlines():
            line = line.strip()
            normalized = "".join(
                char for char in unicodedata.normalize("NFKD", line.casefold())
                if char.isalnum()
            )
            if not line or not normalized or normalized in seen_lines:
                continue
            seen_lines.add(normalized)
            unique_lines.append(line)
        text = "\n".join(unique_lines)
        if len(text) > budget:
            truncated = True
            text = text[:budget].rsplit("\n", 1)[0] or text[:budget]
        blocks.append(f"Source complémentaire — {label} :\n{text}")
    result = "\n\n".join(blocks)
    if truncated:
        result += (
            "\n\nNote : plusieurs sources sont longues ; certaines lignes ont ete omises pour limiter "
            "la taille du contexte. N'en deduis aucun fait absent des extraits."
        )
    return result


def ollama_status():
    try:
        with urlopen("http://127.0.0.1:11434/api/tags", timeout=1) as response:
            models = json.loads(response.read().decode("utf-8")).get("models", [])
        return {"ready": any(model.get("name") == MODEL_NAME for model in models), "model": MODEL_NAME}
    except (OSError, ValueError, TimeoutError):
        return {"ready": False, "model": MODEL_NAME}


def copilot_status():
    return {"sdk_installed": importlib.util.find_spec("copilot") is not None}


EN_WORDS = {
    "the", "and", "with", "for", "you", "your", "our", "will", "are", "of", "to", "in", "we",
    "is", "this", "that", "have", "from", "as", "be", "or", "by", "an", "experience", "team", "work",
}
FR_WORDS = {
    "le", "la", "les", "des", "et", "pour", "avec", "vous", "votre", "nous", "de", "du", "une",
    "un", "dans", "sur", "est", "au", "aux", "que", "qui", "par", "sont", "ce", "ses",
}


def detect_language(text):
    """Retourne 'en' si le texte est majoritairement anglais, sinon 'fr' (valeur par défaut)."""
    words = re.findall(r"[a-zà-ÿ']+", (text or "").lower())
    english = sum(word in EN_WORDS for word in words)
    french = sum(word in FR_WORDS for word in words)
    return "en" if english > french else "fr"


def resolve_language(choice, application):
    if choice in {"fr", "en"}:
        return choice
    return detect_language(f"{application['position']}\n{application['job_description']}")


def english_task(kind):
    current_year = datetime.now().year
    degree_rule = (
        f"It is {current_year}. Any degree whose end year is {current_year} or earlier is COMPLETED: the "
        "candidate holds it. Never write 'currently studying', 'in progress' or 'student' for such a degree; "
        "write 'Master's degree in ...' for each master in the source. "
    )
    if kind == "cv":
        return degree_rule + (
            "You are a resume writer for data jobs. Write the resume IN ENGLISH, targeted to the job offer, "
            "ATS-friendly and faithful to the source. Use the offer's requirements only to choose and order "
            "facts already present in the source resume: a requirement is never proof of experience. "
            "NEVER invent or extrapolate a tool, task, result, figure, degree, date, level, employer or "
            "responsibility. Do not call the candidate an expert unless the source says so. Keep names, job "
            "titles and dates exactly as in the source and translate the descriptions into natural English. "
            "Do NOT include the name, e-mail, phone, address, city or a header job title: the application adds "
            "them. Do not add 'Name:', 'Email:' or 'Location:' lines. "
            "Return only the resume, with no introduction or comment, using these headings when documented: "
            "PROFILE, PROFESSIONAL EXPERIENCE, PROJECTS, EDUCATION, CERTIFICATIONS and SOFTWARE. "
            "Do NOT write a KEY SKILLS, TECHNICAL SKILLS or SOFT SKILLS section: skills must appear inside each "
            "experience. The application adds LANGUAGES and HOBBIES from the profile when provided. "
            "Think like the recruiter for this role in a very competitive market: identify what they would look "
            "for first in the offer (missions, tools, keywords) and surface the matching elements from all "
            "reference sources, merging common skills without repeating them and reusing the offer's wording "
            "when a source supports it. Do not attach a global skill to a job: the tools listed for an "
            "experience must be explicitly linked to that experience in at least one source. "
            "The versions provided are complementary; if their dates or facts conflict, do not pick arbitrarily "
            "and do not invent a correction. "
            "Profile: 2 factual sentences maximum, focused on the target role, no superlatives. "
            "One page maximum: select at most the 3 experiences most relevant to the offer, without dropping a "
            "key experience in favour of a less relevant one. For each experience keep at most 2 short, highly "
            "significant bullet points starting with an action verb, ordered by relevance. Do not invent tasks "
            "and do not repeat the same idea. Quantify (volumes, percentages, deadlines, number of users) ONLY "
            "with figures present in the source or in the verified achievements; without a figure, describe the "
            "mission without claiming a result. End each job with a single line "
            "'Skills: tool1, tool2, tool3' (5 maximum) listing the tools and methods actually used in that job, "
            "taken from the sources, prioritising those requested by the offer. "
            "Certifications: choose at most 4 certifications from the profile list, only if their title or "
            "skills match the missions and tools of this offer. Copy their titles exactly, without diploma "
            "identifier. Invent no certification or skill. Omit the section if none matches; never copy the "
            "whole list. "
            "SOFTWARE: 8 tools maximum, kept only if requested in the offer or directly useful to its missions. "
            "Projects: at most 1 project, only if more relevant than a retained experience. "
            "Education: keep the degrees from the source resume, one per entry, with degree, institution and "
            "dates on separate lines. Never merge several degrees on one line and do not separate them with "
            "semicolons. Maximum 350 words. Plain text, no tables, icons or skill bars."
        )
    return degree_rule + (
        "Write a convincing, personal and natural cover letter IN ENGLISH for the target job. "
        "Never use em or en dashes or double hyphens: rephrase with a comma, a colon or a new sentence. "
        "The body is 170 to 220 words in 4 paragraphs of prose, with no list, numbering, subheading or dashes. "
        "The letter already starts with 'Dear Hiring Manager,': do not produce any salutation, closing "
        "formula, signature, address, date or subject line, because the application adds them. Finish every "
        "sentence and the last paragraph completely. "
        "1) Introduction: start with 'I am delighted to apply for the position of ...' (or 'I am excited "
        "to apply...'), then give your specific motivation for the company's missions or projects, only if "
        "the offer describes them. Avoid generic compliments and empty phrases. Never start with 'I am "
        "writing to', 'Hello' or 'passionate'. "
        "2) Background: present the relevant degree and the data/statistics experience when the sources "
        "establish them. Never compute a number of years from dates: state a duration only if explicitly "
        "given. "
        "3) Fit: choose the 2 or 3 most important requirements of the offer and justify each with a precise "
        "example (mission, tool or result) present in the resume or references. Use each experience and each "
        "result only once; repeat no example, sentence or paragraph. For each requirement, explain explicitly "
        "what this proof would let you bring to the role and its missions: link the skill, the lived example "
        "and the expected concrete contribution. Select only the certifications whose skills match the "
        "offer's requirements; explain how they can support your contribution, without presenting them as "
        "professional experience and without copying diploma identifiers. Cite none if they add nothing. If a "
        "requirement is not proven, do not claim to master it: choose a verified close skill. Cite at most 2 "
        "experiences, only those closest to the target role, and ignore jobs unrelated to the offer. Each "
        "employer is cited only once. "
        "4) Conclusion: in one or two sentences, soberly highlight open-mindedness, agility, proactivity or "
        "adaptability only if the sources justify it. The application adds the interview sentence: do not "
        "write it. Write 'at [employer]' to place an experience. "
        "Invent no fact, figure, degree, length of experience, tool, responsibility, result or company "
        "quality. Do not claim success unless the sources prove it. Do not repeat the resume line by line "
        "and add no thanks or closing formula."
    )


def build_prompt(kind, profile, application, reference_documents=(), style_examples=(), language="fr"):
    profile_text = "\n".join(
        [
            f"Nom : {profile['full_name']}", f"E-mail : {profile['email']}",
            f"Telephone : {profile['phone']}", f"Localisation : {profile['location']}",
            "Sources de reference du candidat (seules sources autorisees pour les faits personnels) :",
            build_reference_context(profile, reference_documents),
        ]
    )
    achievements = profile_list(profile, "achievements")
    if achievements:
        profile_text += (
            "\nRealisations chiffrees verifiees par le candidat (a utiliser en priorite pour quantifier, "
            "telles quelles) :\n" + "\n".join(f"- {item}" for item in achievements)
        )
    certifications = profile_list(profile, "certifications")
    if certifications:
        profile_text += (
            "\nCertifications et competences associees (formations validees par le candidat, "
            "pas des experiences professionnelles) :\n"
            + "\n".join(f"- {item}" for item in certifications)
        )
    offer_text = "\n".join(
        [
            f"Poste : {application['position']}", f"Entreprise : {application['company']}",
            f"Localisation : {application['location']}", "Description de l'offre :",
            application["job_description"][:3500],
        ]
    )
    if kind == "letter":
        tailored_cv = (application["cv_draft"] or "").strip() if "cv_draft" in application.keys() else ""
        if tailored_cv:
            offer_text += (
                "\n\nCV ADAPTE A CETTE OFFRE (deja valide par le candidat ; reste coherent avec lui : "
                "reprends en priorite les experiences, outils et resultats qui y figurent, sans rien ajouter "
                "qui ne soit pas dans les sources de reference) :\n" + tailored_cv[:2500]
            )
    current_year = datetime.now().year
    degree_rule = (
        f"Nous sommes en {current_year}. Toute formation dont l'annee de fin est {current_year} ou anterieure "
        "est TERMINEE : le candidat en est titulaire. N'ecris jamais « actuellement en », « en cours de », "
        "« en train de preparer » ni « etudiant » pour un diplome dont l'annee de fin est atteinte ; "
        "ecris « titulaire d'un Master ... » pour chaque master de la source. "
    )
    if language == "en":
        task = english_task(kind)
    elif kind == "cv":
        task = degree_rule + (
            "Tu es redacteur de CV francais pour les metiers de la data. Produis un CV cible, lisible par un ATS "
            "et fidele a la source. N'utilise les exigences de l'offre que pour choisir et ordonner les faits "
            "deja presents dans le CV source : une exigence n'est jamais une preuve d'experience. "
            "INTERDICTION d'inventer ou d'extrapoler un outil, une mission, un resultat, un chiffre, une "
            "formation, une date, un niveau, un employeur ou une responsabilite. Ne qualifie pas le candidat "
            "d'expert si la source ne le dit pas. Conserve mot pour mot les noms, intitules et dates. "
            "N'inclus PAS le nom, l'e-mail, le telephone, l'adresse, la ville ni un titre de poste en-tete : "
            "l'application les ajoute. N'ajoute pas de ligne 'Nom :', 'E-mail :' ou 'Localisation :'. "
            "Retourne uniquement le CV, sans introduction ni commentaire, avec ces rubriques si elles sont "
            "documentees : PROFIL, EXPERIENCES PROFESSIONNELLES, PROJETS, FORMATION, CERTIFICATIONS et LOGICIELS. "
            "N'ecris PAS de rubrique COMPETENCES CLES, COMPETENCES TECHNIQUES ou APTITUDES : "
            "les competences doivent apparaitre dans chaque experience. L'application ajoute LANGUES et "
            "LOISIRS depuis le profil lorsqu'ils sont renseignes. "
            "Raisonne comme le recruteur de ce poste, dans un marche tres concurrentiel : retiens d'abord ce "
            "qu'il chercherait en priorite dans l'offre (missions, outils, mots-cles) et fais remonter les "
            "elements pertinents de toutes les sources de reference, en fusionnant les competences communes "
            "sans les repeter, et en reprenant les mots de l'offre quand une source les justifie. "
            "Ne rattache pas automatiquement une competence globale a un poste : les outils indiques pour "
            "une experience doivent etre explicitement relies a cette experience dans au moins une source. "
            "Les versions fournies sont complementaires ; si leurs dates ou faits se contredisent, ne choisis "
            "pas arbitrairement et n'invente pas de correction. "
            "Profil : 2 phrases factuelles maximum, centrees sur le poste vise, sans superlatif. "
            "Une page maximum : selectionne au plus les 3 experiences les plus pertinentes pour l'offre, "
            "sans supprimer une experience cle au profit d'un poste moins pertinent. Pour chaque experience, "
            "garde au plus 2 activites courtes et hautement significatives, commencant par un verbe d'action "
            "et classees par pertinence. N'invente pas d'activites et ne repete pas la meme idee. Quantifie (volumes, "
            "pourcentages, delais, nombre d'utilisateurs) UNIQUEMENT avec les chiffres presents dans la "
            "source ou dans les realisations chiffrees ; sans chiffre, decris la mission sans pretendre a "
            "un resultat. Termine chaque poste par une ligne unique "
            "'Competences : outil1, outil2, outil3' (5 maximum) listant les outils et methodes reellement "
            "utilises dans ce poste, tires des sources, en priorite ceux demandes par l'offre. "
            "Certifications : choisis au maximum 4 certifications de la liste du profil, uniquement si leur "
            "intitule ou leurs competences correspondent aux missions et outils de cette offre. Recopie leurs "
            "intitules exactement, sans identifiant de diplome. N'invente ni certification ni competence. "
            "Omet cette rubrique si aucune certification ne correspond ; ne recopie jamais toute la liste. "
            "LOGICIELS : 8 outils maximum, gardes uniquement s'ils sont demandes dans l'offre ou directement "
            "utiles a ses missions. "
            "Projets : au plus 1 projet, uniquement s'il est plus pertinent qu'une experience conservee. "
            "Formation : conserve les diplomes du CV source, une formation par entree, avec "
            "diplome, etablissement et dates sur "
            "des lignes distinctes. Ne fusionne jamais plusieurs diplomes sur une ligne et ne les separe pas "
            "par des points-virgules. Maximum 350 mots. Texte simple, sans tableau, icones ni jauges."
        )
    else:
        task = degree_rule + (
            "Redige une lettre de motivation convaincante, personnelle et naturelle, en francais, pour le poste vise. "
            "N'utilise jamais de tiret long ou demi-cadratin (— ou –) ni de double tiret : reformule avec une "
            "virgule, deux-points ou une nouvelle phrase. "
            "Le corps fait 170 a 220 mots en 4 paragraphes de prose, sans liste, numerotation, sous-titre ni tirets. "
            "La lettre commence deja par « Madame, Monsieur, » : ne produis ni salutation ni formule de politesse, "
            "signature, adresse, date ou objet, car l'application les ajoute. Termine chaque phrase et le dernier "
            "paragraphe completement. "
            "1) Introduction : commence par « C'est avec grand plaisir que je vous adresse ma candidature au poste "
            "de ... » (ou « C'est avec enthousiasme que... »), puis donne ta motivation specifique pour les missions "
            "ou les projets de l'entreprise, uniquement si l'offre les decrit. Evite les compliments generiques "
            "et les formules creuses. Ne commence jamais par « J'ai le plaisir de vous proposer », "
            "« Je suis tres heureux », « Bonjour » ou « passionne ». "
            "2) Parcours : presente le diplome pertinent et l'experience en data/statistiques lorsqu'ils sont "
            "etablis par les sources. Ne calcule jamais un nombre d'annees a partir des dates : annonce une duree "
            "uniquement si elle est explicitement indiquee. "
            "3) Adequation : choisis les 2 ou 3 exigences les plus importantes de l'offre et justifie chacune par "
            "un exemple precis (mission, outil ou resultat) present dans le CV ou les references. N'utilise chaque "
            "experience et chaque resultat qu'une seule fois ; ne repete ni exemple, ni phrase, ni paragraphe. "
            "Pour chaque exigence, explique explicitement ce que cette preuve te permettrait d'apporter au poste "
            "et a ses missions : relie la competence, l'exemple vecu et la contribution concrete attendue. "
            "Selectionne uniquement les certifications dont les competences correspondent aux exigences de l'offre ; "
            "explique en quoi elles peuvent soutenir ta contribution, sans les presenter comme une experience "
            "professionnelle et sans recopier leurs identifiants de diplome. N'en cite aucune si elles n'apportent "
            "rien a cette offre. Si une exigence n'est "
            "pas prouvee, ne pretend pas la maitriser : choisis une competence proche verifiee. "
            "Cite au maximum 2 experiences, uniquement les plus proches du poste vise, et ignore les postes "
            "peu lies a l'offre (par exemple suivi budgetaire ou autre fonction non data science). "
            "Chaque employeur n'est cite qu'une fois. "
            "4) Conclusion : en une ou deux phrases, souligne sobrement l'ouverture d'esprit, l'agilite, la proactivite ou l'adaptation "
            "seulement si les sources le justifient. L'application ajoute la phrase sur l'entretien : ne la redige pas. "
            "Ecris « aupres de [employeur] » pour situer une experience, jamais « chez [employeur] ». "
            "N'invente aucun fait, chiffre, diplome, duree d'experience, outil, responsabilite, resultat ni "
            "qualite de l'entreprise. N'affirme pas que tu as reussi quelque chose si les sources ne le prouvent. "
            "Ne repete pas le CV ligne par ligne et n'ajoute pas de remerciement ou formule de fin."
        )
    examples_text = build_examples_context(kind, style_examples)
    return f"{task}{examples_text}\n\nPROFIL\n{profile_text}\n\nOFFRE\n{offer_text}"


def clean_cv_draft(content):
    cleaned_lines = []
    for line in content.splitlines():
        line = re.sub(r"^\s{0,3}#{1,3}\s*", "", line.strip())
        line = line.replace("**", "").replace("__", "")
        if not line:
            continue
        if re.match(r"(?i)^(nom|name|e-?mail|telephone|t[eé]l[eé]phone|phone|localisation|location|adresse|address)\s*:", line):
            continue
        if not cleaned_lines and re.match(
            r"(?i)^(cv\s+(adapte|adapt\u00e9)|(tailored\s+)?(cv|resume))(\s*[-:\u2013\u2014].*)?$", line
        ):
            continue
        if not cleaned_lines and re.match(r"(?i)^(voici|bien s\u00fbr|certainement|here is|here's|sure|certainly)\b", line):
            continue
        cleaned_lines.append(line)
    return "\n".join(cleaned_lines).strip()


def normalize_dashes(text):
    # Les tirets longs (— – --) trahissent un texte généré : on garde un simple trait d'union.
    text = re.sub(r"\s+(?:—|–|--+)\s+", " - ", text)
    return re.sub(r"(?<=\S)(?:—|–|--+)(?=\S)", "-", text)


def clean_letter_draft(content, profile_name=""):
    cv_headings = {
        "A PROPOS", "PROFIL", "COMPETENCES", "COMPETENCES CLES", "EXPERIENCE",
        "EXPERIENCES", "EXPERIENCES PROFESSIONNELLES", "FORMATION", "PROJETS", "LANGUES",
        "EDUCATION", "EXPERIENCE PROFESSIONNELLE", "COMPETENCES TECHNIQUES",
        "SUMMARY", "PROFESSIONAL EXPERIENCE", "WORK EXPERIENCE", "PROJECTS", "LANGUAGES", "SKILLS",
        "TECHNICAL SKILLS", "KEY SKILLS", "SOFTWARE", "CERTIFICATIONS",
    }
    cleaned_lines = []
    for line in content.splitlines():
        line = re.sub(r"^\s{0,3}#{1,3}\s*", "", line.strip())
        line = line.replace("**", "").replace("__", "").replace("`", "").strip().strip("*").strip()
        if re.fullmatch(r"[-_=*\s]{3,}", line):
            continue
        if not line:
            cleaned_lines.append("")
            continue
        normalized = "".join(
            character
            for character in unicodedata.normalize(
                "NFKD", line.strip().rstrip(":").strip().upper().replace("’", "'")
            )
            if not unicodedata.combining(character)
        )
        if normalized in cv_headings or normalized in {"MOTIVATION LETTER", "LETTRE DE MOTIVATION"}:
            continue
        if re.match(r"^(LETTRE DE MOTIVATION|MOTIVATION LETTER|COVER LETTER)\b", normalized):
            continue
        if normalized.startswith(("OBJET :", "OBJECT:", "SUBJECT:", "RE:")):
            continue
        if normalized.startswith((
            "JE VOUS PRIE D'AGREER", "VEUILLEZ AGREER", "JE VOUS PRIE DE CROIRE",
            "DANS L'ATTENTE DE VOTRE REPONSE", "JE RESTE A VOTRE DISPOSITION",
            "THANK YOU FOR YOUR TIME", "THANK YOU FOR YOUR CONSIDERATION", "I LOOK FORWARD TO HEARING",
        )):
            continue
        if re.fullmatch(
            r"(BONJOUR|BONSOIR|CORDIALEMENT|SINCEREMENT|BIEN CORDIALEMENT|SINCERELY|YOURS SINCERELY|"
            r"YOURS FAITHFULLY|BEST REGARDS|KIND REGARDS|REGARDS|BEST|HELLO)[,.! ]*",
            normalized,
        ):
            continue
        if re.match(r"^DEAR (HIRING|RECRUITMENT|SIR|MADAM|TEAM)\b[^.!?]{0,40}[,:]?$", normalized):
            continue
        if profile_name and normalized == profile_name.strip().upper():
            continue
        if not cleaned_lines and re.match(r"(?i)^(voici|bien s\u00fbr|certainement)\b", line):
            continue
        if normalized in {"MONSIEUR/MADAME,", "MADAME/MONSIEUR,", "MONSIEUR, MADAME,"}:
            line = "Madame, Monsieur,"
        cleaned_lines.append(line)
    cleaned = re.sub(r"\n{3,}", "\n\n", "\n".join(cleaned_lines)).strip()
    cleaned = re.sub(r"(?im)^Madame,\s*Monsieur,\s*$", "", cleaned, count=1).strip()
    cleaned = re.sub(r"(?i)^\s*(?:Madame\s*,?\s*Monsieur|Monsieur\s*,?\s*Madame)\s*,?\s*", "", cleaned, count=1)
    if cleaned[:1].islower():
        cleaned = cleaned[0].upper() + cleaned[1:]
    return normalize_dashes(cleaned)


REFUSAL_PATTERNS = (
    "description complete", "pouvez-vous me fournir", "pouvez vous me fournir", "je dois acceder",
    "je ne peux pas", "informations supplementaires", "extrait fourni", "extrait est insuffisant",
    "as an ai", "i need", "could you provide",
)


def looks_like_refusal(draft):
    plain = "".join(
        character for character in unicodedata.normalize("NFKD", draft.lower().replace("’", "'"))
        if not unicodedata.combining(character)
    )
    return any(pattern in plain for pattern in REFUSAL_PATTERNS)


def generate_with_ollama(prompt, max_tokens, num_ctx=OLLAMA_CONTEXT):
    payload = json.dumps(
        {"model": MODEL_NAME, "messages": [{"role": "user", "content": prompt}],
         "stream": False, "options": {
             "temperature": 0.2, "num_predict": max_tokens, "num_ctx": num_ctx
         }}
    ).encode("utf-8")
    for attempt in range(3):
        req = Request(
            "http://127.0.0.1:11434/api/chat", data=payload,
            headers={"Content-Type": "application/json"}, method="POST"
        )
        try:
            with urlopen(req, timeout=OLLAMA_TIMEOUT) as response:
                result = json.loads(response.read().decode("utf-8"))
            break
        except HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")[:300]
            raise RuntimeError(f"Ollama a repondu avec le code {error.code}: {detail}") from error
        except (URLError, TimeoutError) as error:
            reason = getattr(error, "reason", error)
            connection_refused = isinstance(reason, ConnectionRefusedError)
            app.logger.warning("Echec Ollama, tentative %s/3: %s", attempt + 1, error)
            if connection_refused and attempt < 2:
                time.sleep(attempt + 1)
                continue
            if isinstance(reason, TimeoutError):
                raise RuntimeError(
                    "Qwen n'a pas termine dans le delai imparti. Il tourne sur CPU : "
                    "essaie une annonce plus courte ou relance avec moins de texte."
                ) from error
            raise RuntimeError(
                "Ollama ne repond pas sur 127.0.0.1:11434. Ouvre Ollama, puis reessaie. "
                f"Detail : {reason}"
            ) from error
    content = result.get("message", {}).get("content", "").strip()
    if not content:
        raise RuntimeError("Ollama n'a retourne aucun texte.")
    return content


def generate_with_copilot(prompt, max_tokens):
    try:
        from copilot import CopilotClient
        from copilot.generated.session_events import AssistantMessageData
        from copilot.session import PermissionRequestResult
    except ImportError as error:
        raise RuntimeError(
            "Le SDK Copilot n'est pas installé. Suis les étapes de configuration indiquées dans le README."
        ) from error

    def reject_permissions(_request, _invocation):
        return PermissionRequestResult(kind="reject")

    def generate(working_directory):
        async def run():
            async with CopilotClient() as client:
                async with await client.create_session(
                    available_tools=[],
                    on_permission_request=reject_permissions,
                    enable_config_discovery=False,
                    working_directory=working_directory,
                ) as copilot_session:
                    response = await copilot_session.send_and_wait(
                        prompt, timeout=COPILOT_TIMEOUT
                    )
                    data = getattr(response, "data", None)
                    if not isinstance(data, AssistantMessageData):
                        app.logger.warning(
                            "Copilot returned no assistant message (response data type: %s)",
                            type(data).__name__,
                        )
                        raise RuntimeError(
                            "Copilot a terminé sans fournir de texte. Aucun brouillon n'a été enregistré. "
                            "Vérifie ta connexion GitHub et ton quota Copilot, puis réessaie ou choisis Qwen local."
                        )
                    return (data.content or "").strip()

        return run()

    with tempfile.TemporaryDirectory(prefix="candidatures-copilot-") as working_directory:
        try:
            content = asyncio.run(generate(working_directory))
        except TimeoutError as error:
            raise RuntimeError(
                "Copilot n'a pas terminé à temps. Réessaie avec une offre ou un CV plus court."
            ) from error
        except RuntimeError:
            raise
        except Exception as error:
            app.logger.exception("Échec de la génération avec Copilot SDK")
            raise RuntimeError(
                "Copilot SDK a échoué. Vérifie la connexion GitHub, le quota et le budget de crédits Copilot."
            ) from error
    if not content:
        app.logger.warning("Copilot returned an empty assistant message")
        raise RuntimeError(
            "Copilot a terminé sans fournir de texte. Aucun brouillon n'a été enregistré. "
            "Vérifie ta connexion GitHub et ton quota Copilot, puis réessaie ou choisis Qwen local."
        )
    return content


GENDER_MARKER = re.compile(r"[\s(\[]*\b(?:H\s*/\s*F|F\s*/\s*H|M\s*/\s*F|F\s*/\s*M)(?:\s*/\s*[XN])?\b.*$", re.I)


def cv_role_title(application, upper=True):
    # Les titres de pages d'offres peuvent inclure l'entreprise et la plateforme après le vrai intitulé.
    role = re.sub(
        r"\s*(?:\||[-–—])\s*(?:LinkedInn?|Indeed|HelloWork|APEC)\s*$",
        "",
        application["position"],
        flags=re.I,
    ).strip(" -–—|,()[]")
    company = (application["company"] or "").strip() if "company" in application.keys() else ""
    if len(company) >= 3:
        role = re.sub(
            rf"\s*(?:\||[-–—])\s*{re.escape(company)}\s*$",
            "",
            role,
            flags=re.I,
        ).strip(" -–—|,()[]")
    # Retire la mention H/F et tout ce qui suit (ville).
    role = GENDER_MARKER.sub("", role).strip(" -–—|,()[]")
    # Retire aussi la ville (champ "lieu" de l'offre) quand elle est collée au titre sans H/F.
    location = (application["location"] or "") if "location" in application.keys() else ""
    for city in re.split(r"[,()/\-–—]", location):
        city = city.strip()
        if len(city) >= 3:
            role = re.sub(rf"[\s\-–—|,(]*\b{re.escape(city)}\b[\s)]*$", "", role, flags=re.I).strip(" -–—|,()[]")
    role = role or application["position"].strip()
    return role.upper() if upper else role


def file_slug(text, max_length=None):
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(character for character in text if not unicodedata.combining(character))
    words = [word for word in re.split(r"[^A-Za-z0-9]+", text) if word]
    if max_length:
        kept = []
        for word in words:
            if kept and len("-".join(kept + [word])) > max_length:
                break
            kept.append(word)
        words = kept
    return "-".join(words)


def export_file_name(profile, application, kind, language="fr"):
    # Le nom complet du profil est supposé écrit « Prénom NOM » ; le NOM (en majuscules) passe en premier.
    words = (profile["full_name"] or "").split()
    last = [word for word in words if word.isupper() and len(word) > 1]
    first = [word for word in words if word not in last]
    full = "-".join(file_slug(word) for word in last + first) or "Candidat"
    full = "-".join(part.upper() if part in map(file_slug, last) else part.capitalize() for part in full.split("-"))
    role = file_slug(cv_role_title(application, upper=False), 20)
    letter_code = "CL" if language == "en" else "LM"
    parts = [full, "CV" if kind == "cv" else letter_code, role, file_slug(application["company"], 18)]
    return "_".join(part for part in parts if part)


def cv_location(profile, language="fr"):
    place = (profile["location"] or "").strip()
    mobility = "(Nationwide mobility)" if language == "en" else "(Mobilité Nationale)"
    return f"{place} {mobility}" if place else mobility


def make_docx(title, content):
    document = Document()
    document.add_heading(title, level=0)
    for line in content.splitlines():
        cleaned = line.strip()
        if not cleaned:
            continue
        if cleaned.startswith("### "):
            document.add_heading(cleaned[4:], level=2)
        elif cleaned.startswith("## "):
            document.add_heading(cleaned[3:], level=1)
        elif cleaned.startswith("# "):
            document.add_heading(cleaned[2:], level=1)
        elif cleaned.startswith(("- ", "* ")):
            document.add_paragraph(cleaned[2:], style="List Bullet")
        else:
            document.add_paragraph(cleaned)
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def french_date():
    months = [
        "janvier", "fevrier", "mars", "avril", "mai", "juin",
        "juillet", "aout", "septembre", "octobre", "novembre", "decembre",
    ]
    today = datetime.now()
    return f"{today.day} {months[today.month - 1]} {today.year}"


def english_date():
    today = datetime.now()
    return f"{today.strftime('%B')} {today.day}, {today.year}"


LETTER_TEXT = {
    "fr": {
        "attention": "À l’attention du service recrutement",
        "subject_label": "Objet :",
        "subject": "Candidature au poste de {role}",
        "salutation": "Madame, Monsieur,",
        "closing": "Je vous prie d’agréer, Madame, Monsieur, l’expression de mes salutations distinguées.",
        "interview": (
            "Impatient de vous rencontrer lors d’un entretien, je serai heureux de vous exposer "
            "plus en détail ma motivation et la manière dont je pourrai contribuer à vos projets."
        ),
        "title": "Lettre de motivation",
    },
    "en": {
        "attention": "To the attention of the recruitment team",
        "subject_label": "Subject:",
        "subject": "Application for the position of {role}",
        "salutation": "Dear Hiring Manager,",
        "closing": "Yours sincerely,",
        "interview": (
            "I would welcome the opportunity to discuss my motivation in an interview and explain "
            "how I can contribute to your projects."
        ),
        "title": "Cover letter",
    },
}


def letter_date(profile, language="fr"):
    place = (profile["location"] or "").strip()
    if language == "en":
        return f"{place}, {english_date()}" if place else english_date()
    return f"{place}, le {french_date()}" if place else french_date()


def letter_header_lines(profile, application, location=None):
    contacts = [
        value.strip()
        for value in (profile["location"] if location is None else location, profile["email"], profile["phone"])
        if value and value.strip()
    ]
    return contacts


def letter_body_paragraphs(profile, content, language="fr"):
    cleaned = clean_letter_draft(content, profile["full_name"])
    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(r"\n\s*\n", cleaned)
        if paragraph.strip()
    ]
    # Une generation coupee laisse une phrase inachevee : on la retire.
    while paragraphs and not re.search(r"[.!?…]['»”’)]?$", paragraphs[-1]):
        complete = re.match(r"(?s)^(.*[.!?…])\s", paragraphs[-1] + " ")
        if complete and len(complete.group(1)) > 40:
            paragraphs[-1] = complete.group(1).strip()
            break
        paragraphs.pop()
    if paragraphs:
        paragraphs.append(LETTER_TEXT[language]["interview"])
    return paragraphs


def make_letter_docx(profile, application, content):
    language = detect_language(content)
    text = LETTER_TEXT[language]
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = 11
    normal.paragraph_format.space_after = 9

    name = document.add_paragraph()
    name.paragraph_format.space_after = 2
    name_run = name.add_run(profile["full_name"].strip() or "Nom a completer")
    name_run.bold = True
    name_run.font.name = "Aptos Display"
    name_run.font.size = 16
    if profile["postal_address"].strip():
        address = document.add_paragraph(profile["postal_address"].strip())
        address.paragraph_format.space_after = 2
    contacts = letter_header_lines(profile, application)
    if contacts:
        contact = document.add_paragraph("  |  ".join(contacts))
        contact.paragraph_format.space_after = 14

    recipient_lines = [application["company"].strip()]
    recipient_lines.extend(
        line.strip()
        for line in (application["company_address"] or "").splitlines()
        if line.strip()
    )
    for line in recipient_lines:
        recipient = document.add_paragraph(line)
        recipient.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        recipient.paragraph_format.space_after = 1

    date_line = document.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_line.add_run(letter_date(profile, language))
    attention = document.add_paragraph(text["attention"])
    attention.paragraph_format.space_before = 8
    subject = document.add_paragraph()
    subject.paragraph_format.space_before = 8
    subject.paragraph_format.space_after = 14
    subject_run = subject.add_run(
        f"{text['subject_label']} {text['subject'].format(role=cv_role_title(application, upper=False))}"
    )
    subject_run.bold = True

    document.add_paragraph(text["salutation"])
    for paragraph_text in letter_body_paragraphs(profile, content, language):
        document.add_paragraph(paragraph_text)
    document.add_paragraph(text["closing"])
    document.add_paragraph(profile["full_name"].strip() or "[Nom a completer]")

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def make_cv_docx(profile, application, content):
    language = detect_language(content)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = 10
    normal.paragraph_format.space_after = 4

    name = profile["full_name"].strip() or "Nom a completer"
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = 2
    run = heading.add_run(name)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = 21

    role_line = document.add_paragraph()
    role_line.paragraph_format.space_after = 5
    role = role_line.add_run(cv_role_title(application))
    role.bold = True
    role.font.size = 11
    role.font.color.rgb = RGBColor(33, 73, 63)

    contacts = [
        value.strip()
        for value in (cv_location(profile, language), profile["email"], profile["phone"])
        if value and value.strip()
    ]
    if contacts:
        contact_line = document.add_paragraph("  |  ".join(contacts))
        contact_line.paragraph_format.space_after = 10

    headings = CV_DOC_HEADINGS
    current_section = ""
    for line in clean_cv_draft(content).splitlines():
        text = line.strip()
        plain = text.strip("*# ").rstrip(":").strip()
        if plain.upper() in headings:
            current_section = CV_SECTION_ALIASES.get(normalize_heading(plain), normalize_heading(plain))
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = 8
            paragraph.paragraph_format.space_after = 3
            run = paragraph.add_run(plain.upper())
            run.bold = True
            run.font.size = 11
            run.font.color.rgb = RGBColor(33, 73, 63)
        elif text.startswith(("- ", "* ", "• ")):
            document.add_paragraph(text[2:].strip(), style="List Bullet")
        elif current_section == "EXPERIENCES PROFESSIONNELLES" and re.search(r"[.!?]$", text):
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("CandidateName", parent=base["Title"], fontName="Helvetica-Bold", fontSize=19, leading=22, textColor=colors.HexColor("#202922"), alignment=0, spaceAfter=3),
        "role": ParagraphStyle("CandidateRole", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=colors.HexColor("#21493f"), spaceAfter=5),
        "contact": ParagraphStyle("CandidateContact", parent=base["Normal"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#59645c"), spaceAfter=10),
        "section": ParagraphStyle("CVSection", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=colors.HexColor("#21493f"), spaceBefore=8, spaceAfter=3, keepWithNext=True),
        "body": ParagraphStyle("DocumentBody", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor("#202922"), spaceAfter=7),
        "date": ParagraphStyle("LetterDate", parent=base["Normal"], fontName="Helvetica", fontSize=9, leading=12, alignment=TA_RIGHT, spaceBefore=10, spaceAfter=12),
        "recipient": ParagraphStyle("LetterRecipient", parent=base["Normal"], fontName="Helvetica", fontSize=9.5, leading=12, alignment=TA_RIGHT, spaceAfter=1),
        "subject": ParagraphStyle("LetterSubject", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=8, spaceAfter=14),
    }


def make_pdf(profile, application, content, kind):
    language = detect_language(content)
    text_labels = LETTER_TEXT[language]
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=2.1 * cm, leftMargin=2.1 * cm,
        topMargin=1.7 * cm, bottomMargin=1.7 * cm,
        title=f"{'CV' if kind == 'cv' else text_labels['title']} - {application['company']}",
    )
    styles = pdf_styles()
    story = []
    name = escape(profile["full_name"].strip() or "Nom a completer")
    header_location = cv_location(profile, language) if kind == "cv" else None
    contacts = "  |  ".join(escape(value) for value in letter_header_lines(profile, application, header_location))

    story.append(Paragraph(name, styles["name"]))
    if kind == "cv":
        story.append(Paragraph(escape(cv_role_title(application)), styles["role"]))
        if contacts:
            story.append(Paragraph(contacts, styles["contact"]))
        content = clean_cv_draft(content)
        headings = CV_DOC_HEADINGS
    else:
        if profile["postal_address"].strip():
            for line in profile["postal_address"].splitlines():
                if line.strip():
                    story.append(Paragraph(escape(line.strip()), styles["contact"]))
        if contacts:
            story.append(Paragraph(contacts, styles["contact"]))
        story.append(Spacer(1, 12))
        story.append(Paragraph(escape(application["company"]), styles["recipient"]))
        for line in (application["company_address"] or "").splitlines():
            if line.strip():
                story.append(Paragraph(escape(line.strip()), styles["recipient"]))
        date_text = escape(letter_date(profile, language))
        story.append(Paragraph(date_text, styles["date"]))
        story.append(Paragraph(text_labels["attention"], styles["body"]))
        role_text = escape(cv_role_title(application, upper=False))
        story.append(Paragraph(
            f"{text_labels['subject_label']} {text_labels['subject'].format(role=role_text)}", styles["subject"]
        ))
        content = clean_letter_draft(content)
        headings = set()

    if kind == "letter":
        story.append(Paragraph(text_labels["salutation"], styles["body"]))
        for paragraph in letter_body_paragraphs(profile, content, language):
            story.append(Paragraph(escape(paragraph), styles["body"]))
        story.append(Paragraph(text_labels["closing"], styles["body"]))
        story.append(Spacer(1, 8))
        story.append(Paragraph(name, styles["body"]))

    for line in content.splitlines() if kind == "cv" else []:
        text = line.strip()
        if not text:
            story.append(Spacer(1, 3))
            continue
        plain = text.strip("*# ").rstrip(":").strip()
        if kind == "cv" and plain.upper() in headings:
            story.append(Paragraph(escape(plain.upper()), styles["section"]))
        elif text.startswith(("- ", "* ", "• ")):
            story.append(Paragraph("&bull; " + escape(text[2:].strip()), styles["body"]))
        else:
            story.append(Paragraph(escape(text.replace("**", "")), styles["body"]))
    if kind == "cv":
        document.build([
            KeepInFrame(document.width, document.height, story, mode="shrink")
        ])
    else:
        document.build(story)
    output.seek(0)
    return output


LATEX_SPECIAL = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
    "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}
CV_SECTION_TITLES = {
    "PROFIL": "Profil", "COMPETENCES CLES": "Compétences clés",
    "EXPERIENCES PROFESSIONNELLES": "Expériences professionnelles",
    "FORMATION": "Formation", "PROJETS": "Projets", "LANGUES": "Langues",
    "APTITUDES": "Aptitudes", "LOGICIELS": "Outils",
    "COMPETENCES TECHNIQUES": "Compétences techniques",
    "CERTIFICATIONS": "Certifications", "LOISIRS": "Loisirs",
}
CV_SECTION_TITLES_EN = {
    "PROFIL": "Profile", "COMPETENCES CLES": "Key skills",
    "EXPERIENCES PROFESSIONNELLES": "Professional experience",
    "FORMATION": "Education", "PROJETS": "Projects", "LANGUES": "Languages",
    "APTITUDES": "Skills", "LOGICIELS": "Tools",
    "COMPETENCES TECHNIQUES": "Technical skills",
    "CERTIFICATIONS": "Certifications", "LOISIRS": "Hobbies",
}
CV_SECTION_ALIASES = {
    "CERTIFICATION": "CERTIFICATIONS", "OUTILS": "LOGICIELS", "LOISIR": "LOISIRS",
    "CENTRES D'INTERET": "LOISIRS",
    "PROFILE": "PROFIL", "SUMMARY": "PROFIL", "PROFESSIONAL SUMMARY": "PROFIL",
    "PROFESSIONAL EXPERIENCE": "EXPERIENCES PROFESSIONNELLES",
    "WORK EXPERIENCE": "EXPERIENCES PROFESSIONNELLES", "EXPERIENCE": "EXPERIENCES PROFESSIONNELLES",
    "EXPERIENCES": "EXPERIENCES PROFESSIONNELLES",
    "EDUCATION": "FORMATION", "PROJECTS": "PROJETS", "LANGUAGES": "LANGUES",
    "SOFTWARE": "LOGICIELS", "TOOLS": "LOGICIELS", "KEY SKILLS": "COMPETENCES CLES",
    "TECHNICAL SKILLS": "COMPETENCES TECHNIQUES", "SKILLS": "COMPETENCES CLES",
    "HOBBIES": "LOISIRS", "INTERESTS": "LOISIRS",
}
CV_DOC_HEADINGS = {
    "PROFIL", "COMPETENCES CLES", "COMPÉTENCES CLÉS",
    "EXPERIENCES PROFESSIONNELLES", "EXPÉRIENCES PROFESSIONNELLES",
    "FORMATION", "PROJETS", "LANGUES",
    "PROFILE", "SUMMARY", "PROFESSIONAL EXPERIENCE", "WORK EXPERIENCE", "EDUCATION",
    "PROJECTS", "LANGUAGES", "KEY SKILLS",
}
CV_MAIN_SECTIONS = [
    "PROFIL", "EXPERIENCES PROFESSIONNELLES", "PROJETS", "FORMATION", "CERTIFICATIONS",
]
CV_SKILL_SECTIONS = [
    "COMPETENCES CLES", "COMPETENCES TECHNIQUES", "LOGICIELS", "APTITUDES", "LANGUES",
]
LATEX_CV_TEMPLATE = r"""\documentclass[10pt,a4paper]{article}
\usepackage[a4paper,margin=@@MARGIN@@cm,top=@@VMARGIN@@cm,bottom=@@VMARGIN@@cm]{geometry}
\usepackage{fontspec}
\usepackage{xcolor}
\usepackage{graphicx}
\usepackage[hidelinks]{hyperref}
\setmainfont{Latin Modern Roman}
\definecolor{accent}{HTML}{@@ACCENT@@}
\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{2pt}
\newenvironment{cvitems}{\begin{itemize}\setlength{\itemsep}{0.5pt}\setlength{\topsep}{1pt}\setlength{\parsep}{0pt}\setlength{\leftmargin}{1.2em}}{\end{itemize}}
\newcommand{\cvsection}[1]{\vspace{@@GAP2@@pt}{\color{accent}\bfseries\fontsize{@@H@@}{@@HL@@}\selectfont\MakeUppercase{#1}}\par\vspace{-4pt}{\color{accent}\rule{\linewidth}{0.6pt}}\par\vspace{2pt}}
\newcommand{\cventry}[3]{\vspace{@@GAP@@pt}{\bfseries\fontsize{@@T@@}{@@TL@@}\selectfont #1}\hfill{\color{black!65}\fontsize{@@S@@}{@@SL@@}\selectfont #3}\par\vspace{1pt}{\color{accent}\itshape\fontsize{@@S@@}{@@SL@@}\selectfont #2}\par\vspace{@@GAP@@pt}}
\newcommand{\cvskills}[1]{{\color{accent}\bfseries\fontsize{@@S@@}{@@SL@@}\selectfont @@SKILLSLABEL@@} {\fontsize{@@S@@}{@@SL@@}\selectfont #1}\par}
\newsavebox{\cvcontent}
\begin{document}
\sbox{\cvcontent}{%
\begin{minipage}{\textwidth}
{\fontsize{22}{25}\selectfont\bfseries @@NAME@@}\par
{\color{accent}\bfseries\fontsize{10.5}{12}\selectfont @@ROLE@@}\par
{\small @@CONTACT@@}\par
\fontsize{@@F@@}{@@FL@@}\selectfont
@@BODY@@
\end{minipage}}
\ifdim\dimexpr\ht\cvcontent+\dp\cvcontent\relax>0.96\textheight
  \resizebox{!}{0.96\textheight}{\usebox{\cvcontent}}
\else
  \usebox{\cvcontent}
\fi
\end{document}
"""


CV_LAYOUT_DEFAULTS = {"font": 8.5, "margin": 1.4, "spacing": 1.0, "accent": "21493F"}
CV_ACCENTS = {"21493F": "Vert", "1F3A5F": "Bleu", "5A1F2B": "Bordeaux", "222222": "Noir"}


def parse_cv_layout(raw=None):
    """Valide les reglages de mise en page ; toute valeur invalide revient au defaut."""
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except ValueError:
            raw = {}
    raw = raw if isinstance(raw, dict) else {}
    layout = dict(CV_LAYOUT_DEFAULTS)
    for key, low, high in (("font", 7.0, 11.0), ("margin", 0.8, 2.5), ("spacing", 0.5, 2.0)):
        try:
            value = float(raw.get(key, layout[key]))
        except (TypeError, ValueError):
            continue
        if low <= value <= high:
            layout[key] = round(value, 2)
    if str(raw.get("accent", "")).upper() in CV_ACCENTS:
        layout["accent"] = str(raw["accent"]).upper()
    return layout


def latex_escape(text):
    return "".join(LATEX_SPECIAL.get(character, character) for character in text)


def normalize_heading(text):
    stripped = "".join(
        character for character in unicodedata.normalize("NFKD", text.upper())
        if not unicodedata.combining(character)
    )
    return stripped.strip("*# ").rstrip(":").strip()


def latex_itemize(items):
    if not items:
        return []
    return [r"\begin{cvitems}"] + [rf"\item {latex_escape(item)}" for item in items] + [r"\end{cvitems}"]


CV_DATE_RANGE = re.compile(
    r"(?P<date>(?:(?:0?[1-9]|1[0-2])/\d{4}|\d{4})\s*[–—-]\s*"
    r"(?:(?:0?[1-9]|1[0-2])/\d{4}|\d{4}|present|current|en cours))$",
    re.IGNORECASE,
)


def split_trailing_date(text):
    match = CV_DATE_RANGE.search(text.strip())
    if not match:
        return text.strip(), ""
    return text[:match.start()].strip(" ;,-"), match.group("date").strip()


def select_relevant_items(items, application, generated=(), limit=4):
    candidates = list(items)
    if not candidates:
        return []

    def normalized(value):
        return re.sub(r"[^a-z0-9]", "", normalize_heading(value).casefold())

    candidate_by_key = {normalized(item): item for item in candidates}
    chosen = []
    for item in generated:
        source_item = candidate_by_key.get(normalized(item.lstrip("-*• ").strip()))
        if source_item and source_item not in chosen:
            chosen.append(source_item)
    if chosen:
        return chosen[:limit]

    stop_words = {
        "and", "the", "with", "for", "from", "dans", "pour", "avec", "des",
        "les", "une", "sur", "par", "and", "learn", "learning", "certified",
        "certification", "foundations", "foundation", "data", "science",
    }

    def tokens(value):
        normalized_text = normalize_heading(value).casefold()
        return {
            token for token in re.findall(r"[a-z0-9+#]+", normalized_text)
            if token not in stop_words and (len(token) > 2 or token in {"ai", "bi", "ml", "r"})
        }

    offer_text = " ".join(
        str(application[key] or "")
        for key in ("position", "job_description")
        if key in application.keys()
    )
    offer_tokens = tokens(offer_text)
    aliases = (
        (("pl-300", "power bi"), ("dax", "power bi", "dashboard", "reporting", "analyst")),
        (("power query",), ("power bi", "dax", "dashboard", "reporting")),
        (("dax",), ("power bi", "dashboard", "reporting")),
        (("reinforcement learning",), ("machine learning",)),
        (("deep learning",), ("machine learning",)),
        (("jupyter",), ("python",)),
        (("azure devops",), ("ci/cd", "pipelines")),
        (("generative ai", "llm"), ("large language model",)),
        (("docker",), ("containers", "kubernetes")),
        (("git",), ("version control",)),
    )

    def relevance(item):
        item_tokens = tokens(item)
        score = len(item_tokens & offer_tokens)
        normalized_item = normalize_heading(item).casefold()
        normalized_offer = normalize_heading(offer_text).casefold()
        for triggers, matches in aliases:
            if any(trigger in normalized_item for trigger in triggers) and any(
                match in normalized_offer for match in matches
            ):
                score += 2
        return score

    ranked = sorted(
        enumerate(candidates), key=lambda pair: (-relevance(pair[1]), pair[0])
    )
    return [item for _, item in ranked if relevance(item) > 0][:limit]


CV_SKILLS_LINE = re.compile(
    r"(?i)^[-*•\s]*(?:comp[ée]tences|skills|outils|tools|technologies)(?:\s+cl[ée]s)?\s*:\s*(.+)$"
)


def latex_entries(lines, max_items=None, max_entries=None):
    output, header, items, skills = [], [], [], []
    rendered_entries = 0

    def flush():
        nonlocal rendered_entries
        render_entry = not header or max_entries is None or rendered_entries < max_entries
        if header:
            if render_entry:
                title, dates = split_trailing_date(header[0])
                organisation = header[1] if len(header) > 1 else ""
                dates = " ; ".join([dates, *header[2:]]).strip(" ;")
                output.append(
                    rf"\cventry{{{latex_escape(title)}}}{{{latex_escape(organisation)}}}{{{latex_escape(dates)}}}"
                )
            rendered_entries += 1
        if render_entry:
            output.extend(latex_itemize(items[:max_items] if max_items is not None else items))
        if render_entry and skills:
            output.append(rf"\cvskills{{{latex_escape(', '.join(skills))}}}")
        header.clear()
        items.clear()
        skills.clear()

    for line in lines:
        skills_match = CV_SKILLS_LINE.match(line)
        if skills_match:
            skills.append(skills_match.group(1).strip(" ."))
            continue
        is_bullet = line.startswith(("- ", "* ", "• "))
        text = line[2:].strip() if is_bullet else line
        if is_bullet:
            items.append(text)
        elif len(text) <= 90 and not text.endswith("."):
            if items or skills:
                flush()
            header.append(text)
        else:
            items.append(text)
    flush()
    return output


def latex_education_entries(lines):
    segments = [
        segment.strip().strip(" ;")
        for line in lines
        for segment in line.split(";")
        if segment.strip().strip(" ;")
    ]
    output = []
    pending_title = ""
    pending_institution = ""

    def flush(dates=""):
        nonlocal pending_title, pending_institution
        if pending_title:
            output.append(
                rf"\cventry{{{latex_escape(pending_title)}}}"
                rf"{{{latex_escape(pending_institution)}}}{{{latex_escape(dates)}}}"
            )
        pending_title = ""
        pending_institution = ""

    for segment in segments:
        before_date, dates = split_trailing_date(segment)
        if dates:
            if before_date and pending_title and not pending_institution and re.search(
                r"(?i)universit|institut|[eé]cole|facult[eé]|\bcollege\b|\bschool\b", before_date
            ) and not re.match(r"(?i)^(master|licence|bachelor|doctorat|dipl[oô]me|ing[eé]nieur)", before_date):
                pending_institution = before_date
                flush(dates)
            elif before_date:
                if pending_title:
                    flush()
                pending_title = before_date
                flush(dates)
            elif pending_title:
                flush(dates)
            continue
        if re.fullmatch(r"(?:(?:0?[1-9]|1[0-2])/\d{4}|\d{4})", segment):
            if pending_title:
                flush(segment)
            continue
        is_institution = bool(
            re.match(
                r"(?i)^(universit[eé]s?\b|university\b|institut\b|"
                r"ecole\b|école\b|nantes universit[eé]s\b)",
                segment,
            )
            or re.search(r"(?i)\buniversit[eé]s?\b", segment)
            and not re.match(r"(?i)^(master|licence|bachelor|doctorat)", segment)
        )
        starts_degree = bool(
            re.match(
                r"(?i)^(master|licence|bachelor|doctorat|dipl[oô]me|"
                r"ing[eé]nieur|bts\b|dut\b|but\b)",
                segment,
            )
        )
        if pending_title and is_institution and not pending_institution:
            pending_institution = segment
        elif starts_degree or pending_title:
            flush()
            pending_title = segment
        else:
            pending_title = segment
    flush()
    return output


def make_cv_latex(profile, application, content, layout=None):
    layout = parse_cv_layout(layout)
    language = detect_language(content)
    section_titles = CV_SECTION_TITLES_EN if language == "en" else CV_SECTION_TITLES
    font, spacing = layout["font"], layout["spacing"]
    sections, current = [], None
    for raw_line in clean_cv_draft(content).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        parts = line.split(";") if current and current[0] == "FORMATION" else [line]
        for part in parts:
            part = part.strip()
            if not part:
                continue
            key = CV_SECTION_ALIASES.get(normalize_heading(part), normalize_heading(part))
            if key in CV_SECTION_TITLES:
                current = (key, [])
                sections.append(current)
            else:
                if current is None:
                    current = ("PROFIL", [])
                    sections.append(current)
                current[1].append(part)

    generated_certifications = [
        line
        for key, lines in sections if key == "CERTIFICATIONS"
        for line in lines
    ]
    certifications = select_relevant_items(
        [
            item for item in profile_list(profile, "certifications")
            if not re.match(r"(?i)^identifiant(?: du dipl[oô]me)?\s*:", item)
        ],
        application,
        generated_certifications,
    )
    overrides = {
        "LANGUES": profile_list(profile, "languages"),
        "LOISIRS": profile_list(profile, "hobbies"),
    }
    sections = [
        section for section in sections
        if section[0] != "CERTIFICATIONS" and not overrides.get(section[0])
    ]
    sections += [(key, lines) for key, lines in overrides.items() if lines]
    if certifications:
        sections.append(("CERTIFICATIONS", certifications))

    blocks = {}
    for key, lines in sections:
        if key in {"COMPETENCES CLES", "COMPETENCES TECHNIQUES", "APTITUDES"}:
            continue
        body = [rf"\cvsection{{{latex_escape(section_titles[key])}}}"]
        if key == "FORMATION":
            body.extend(latex_education_entries(lines))
        elif key == "EXPERIENCES PROFESSIONNELLES":
            body.extend(latex_entries(lines, max_items=2, max_entries=3))
        elif key == "PROJETS":
            body.extend(latex_entries(lines, max_items=2, max_entries=1))
        elif key == "LOGICIELS":
            tools, seen = [], set()
            for line in lines:
                for tool in re.split(r"[;,•·]\s*(?![^()]*\))", line.lstrip("-*• ")):
                    tool = tool.strip(" .-*")
                    if tool and tool.casefold() not in seen:
                        seen.add(tool.casefold())
                        tools.append(tool)
            tools = select_relevant_items(tools, application, limit=8)
            body.append(latex_escape(" · ".join(tools)) + r"\par")
        elif key == "CERTIFICATIONS":
            compact_items = " · ".join(
                latex_escape(line.lstrip("-*• ").strip()) for line in lines
            )
            body.append(r"{\fontsize{7.5}{8.2}\selectfont " + compact_items + r"\par}")
        elif key in CV_SKILL_SECTIONS or key == "LOISIRS":
            body.extend(latex_itemize([line.lstrip("-*• ").strip() for line in lines]))
        else:
            for line in lines:
                body.append(latex_escape(line.lstrip("-*• ").strip()) + r"\par")
        blocks.setdefault(key, []).extend(body)

    body_parts = [line for key in CV_MAIN_SECTIONS for line in blocks.get(key, [])]

    def column_row(left_keys, right_keys):
        left = [line for key in left_keys for line in blocks.get(key, [])]
        right = [line for key in right_keys for line in blocks.get(key, [])]
        if left and right:
            return [
                r"\vspace{2pt}\noindent\begin{minipage}[t]{0.56\textwidth}\vspace{0pt}", *left,
                r"\end{minipage}\hfill\begin{minipage}[t]{0.40\textwidth}\vspace{0pt}", *right,
                r"\end{minipage}",
            ]
        return left + right

    body_parts += column_row(("COMPETENCES CLES", "COMPETENCES TECHNIQUES", "APTITUDES"), ("LOGICIELS",))
    body_parts += column_row(("LANGUES",), ("LOISIRS",))
    contacts = []
    for value in (cv_location(profile, language), profile["email"], profile["phone"]):
        value = (value or "").strip()
        if not value:
            continue
        contacts.append(
            rf"\href{{mailto:{value}}}{{{latex_escape(value)}}}"
            if value == (profile["email"] or "").strip() and re.fullmatch(r"[^\s{}\\%#&~^$_]+@[^\s{}\\%#&~^$_]+", value)
            else latex_escape(value)
        )
    for key, label in (("linkedin_url", "LinkedIn"), ("portfolio_url", "Portfolio")):
        if key in profile.keys() and profile[key]:
            link = normalize_profile_url(profile[key])
            contacts.append(rf"\href{{{latex_escape(link)}}}{{{label}}}")
    replacements = {
        "@@NAME@@": latex_escape(profile["full_name"].strip() or "Nom à compléter"),
        "@@ROLE@@": latex_escape(cv_role_title(application)),
        "@@CONTACT@@": r" \;|\; ".join(contacts),
        "@@BODY@@": "\n".join(body_parts),
        "@@MARGIN@@": str(layout["margin"]),
        "@@VMARGIN@@": str(round(layout["margin"] * 0.7, 2)),
        "@@ACCENT@@": layout["accent"],
        "@@SKILLSLABEL@@": "Skills:" if language == "en" else "Compétences :",
        "@@GAP2@@": f"{6 * spacing:.1f}",
        "@@GAP@@": f"{2.5 * spacing:.1f}",
        "@@H@@": f"{font + 2:.1f}", "@@HL@@": f"{font + 3.5:.1f}",
        "@@T@@": f"{font + 1:.1f}", "@@TL@@": f"{font + 2.5:.1f}",
        "@@S@@": f"{font:.1f}", "@@SL@@": f"{font + 1.5:.1f}",
        "@@F@@": f"{font:.1f}", "@@FL@@": f"{font + 1.5 + (spacing - 1) * 1.5:.1f}",
    }
    document = LATEX_CV_TEMPLATE
    for placeholder, value in replacements.items():
        document = document.replace(placeholder, value)
    return document


def make_letter_latex(profile, application, content):
    esc = latex_escape
    language = detect_language(content)
    text = LETTER_TEXT[language]
    name = profile["full_name"].strip() or "Nom à compléter"
    sender = [rf"{{\bfseries\large {esc(name)}}}"]
    sender += [esc(line.strip()) for line in (profile["postal_address"] or "").splitlines() if line.strip()]
    contacts = [(profile[key] or "").strip() for key in ("location", "email", "phone")]
    contacts = [esc(value) for value in contacts if value]
    if contacts:
        sender.append(" | ".join(contacts))
    recipient = [
        esc(text["attention"]),
        rf"{{\bfseries {esc(application['company'])}}}",
    ]
    recipient += [esc(line.strip()) for line in (application["company_address"] or "").splitlines() if line.strip()]
    date_text = letter_date(profile, language)
    paragraphs = "\n\n".join(esc(p) for p in letter_body_paragraphs(profile, content, language))
    return "\n".join([
        r"\documentclass[11pt,a4paper]{article}",
        r"\usepackage[a4paper,margin=2.2cm]{geometry}",
        r"\usepackage{fontspec}",
        r"\setmainfont{Latin Modern Roman}",
        r"\pagestyle{empty}",
        r"\setlength{\parindent}{0pt}",
        r"\setlength{\parskip}{8pt}",
        r"\begin{document}",
        r"\noindent\begin{minipage}[t]{0.7\textwidth}\raggedright",
        r"\\".join(sender),
        r"\end{minipage}\par",
        r"\vspace{28pt}",
        r"\hfill\begin{minipage}[t]{0.42\textwidth}\raggedright",
        r"\\".join(recipient),
        r"\end{minipage}\par",
        r"\vspace{14pt}",
        r"\begin{flushright}" + esc(date_text) + r"\end{flushright}",
        r"\vspace{6pt}",
        rf"\noindent\textbf{{{esc(text['subject_label'])}}} {esc(text['subject'].format(role=cv_role_title(application, upper=False)))}\par",
        r"\vspace{6pt}",
        esc(text["salutation"]) + "\\par",
        paragraphs + r"\par",
        esc(text["closing"]) + "\\par",
        r"\vspace{14pt}",
        r"\begin{flushright}\bfseries " + esc(name) + r"\end{flushright}",
        r"\end{document}",
    ])


def compile_latex(source, one_page=False):
    engine = shutil.which("xelatex")
    if engine is None:
        raise RuntimeError("LaTeX (xelatex) n'est pas installe sur cet ordinateur.")
    with tempfile.TemporaryDirectory() as work_dir:
        (Path(work_dir) / "cv.tex").write_text(source, encoding="utf-8")
        try:
            subprocess.run(
                [engine, "-interaction=nonstopmode", "-halt-on-error", "-no-shell-escape", "cv.tex"],
                cwd=work_dir, capture_output=True, timeout=120, check=True,
            )
        except subprocess.TimeoutExpired as error:
            raise RuntimeError("La compilation LaTeX a depasse le delai.") from error
        except subprocess.CalledProcessError as error:
            log = (Path(work_dir) / "cv.log")
            text = log.read_text(encoding="utf-8", errors="replace") if log.exists() else ""
            detail = next((line for line in text.splitlines() if line.startswith("!")), "")
            raise RuntimeError(f"Echec de la compilation LaTeX. {detail}") from error
        pdf = io.BytesIO((Path(work_dir) / "cv.pdf").read_bytes())
        try:
            reader = PdfReader(pdf, strict=True)
            pages = reader.pages
            text = "\n".join(page.extract_text() or "" for page in pages).strip()
        except (PdfReadError, OSError, ValueError) as error:
            raise RuntimeError("Le PDF compilé est illisible. Essaie l’export PDF de secours.") from error
        if not pages or not text:
            raise RuntimeError("Le PDF compilé ne contient aucune page lisible.")
        if one_page and len(pages) != 1:
            raise RuntimeError(
                "Le PDF du CV dépasse une page. Réduis le contenu puis régénère le CV."
            )
        pdf.seek(0)
        return pdf


@app.teardown_appcontext
def teardown_db(_error):
    close_db(_error)


@app.before_request
def protect_post_requests():
    if request.method == "POST" and request.endpoint != "extension_offer":
        expected = session.get("csrf_token", "")
        provided = request.form.get("csrf_token", "")
        if not expected or not hmac.compare_digest(expected, provided):
            abort(400)


@app.context_processor
def inject_csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["csrf_token"]}


@app.get("/")
def index():
    applications = get_db().execute(
        "SELECT * FROM applications ORDER BY updated_at DESC, id DESC"
    ).fetchall()
    selected_id = request.args.get("selected", type=int)
    selected = get_application(selected_id) if selected_id else None
    if selected is None and applications:
        selected = applications[0]
    profile = get_profile()
    reference_documents = get_reference_documents()
    has_resume = bool(profile["resume_source"].strip() or reference_documents)
    ollama = ollama_status()
    copilot = copilot_status()
    return render_template(
        "index.html", profile=profile, applications=applications,
        reference_documents=reference_documents,
        selected=selected, statuses=STATUSES, ollama=ollama,         copilot=copilot, gmail=job_import.gmail_settings(),
        extension_token=get_extension_token(),
        style_examples=get_style_examples(),
        cv_layout=parse_cv_layout(selected["cv_layout"] if selected else None),
        cv_accents=CV_ACCENTS,
        has_resume=has_resume,
        has_job_description=bool(selected and selected["job_description"].strip()),
        can_generate=bool(has_resume and selected and selected["job_description"].strip()),
        counts={
            "total": len(applications),
            "sent": sum(item["status"] == "Candidature envoyee" for item in applications),
            "interviews": sum(item["status"] == "Entretien" for item in applications),
        },
    )


@app.post("/profile")
def save_profile():
    resume_source = request.form.get("resume_source", "").strip()
    resume_pdf = request.files.get("resume_pdf")
    pdf_uploaded = bool(resume_pdf and resume_pdf.filename)
    if pdf_uploaded:
        if Path(resume_pdf.filename).suffix.lower() != ".pdf":
            flash("Choisis un fichier PDF pour importer ton CV.", "error")
            return redirect(url_for("index", selected=request.form.get("selected", type=int)))
        pdf_content = resume_pdf.read()
        if not pdf_content.startswith(b"%PDF-"):
            flash("Le fichier choisi ne semble pas etre un PDF valide.", "error")
            return redirect(url_for("index", selected=request.form.get("selected", type=int)))
        try:
            reader = PdfReader(io.BytesIO(pdf_content))
            resume_source = "\n\n".join(
                page.extract_text() or "" for page in reader.pages
            ).strip()
        except Exception:
            flash("Le PDF n'a pas pu etre lu. Verifie le fichier puis reessaie.", "error")
            return redirect(url_for("index", selected=request.form.get("selected", type=int)))
        if not resume_source:
            flash(
                "Aucun texte n'a ete extrait. Le PDF est peut-etre scanne ; utilise un PDF "
                "avec du texte selectionnable ou colle le texte du CV.",
                "error",
            )
            return redirect(url_for("index", selected=request.form.get("selected", type=int)))

    try:
        profile_urls = {
            field: normalize_profile_url(request.form.get(field, ""))
            for field in PROFILE_URL_FIELDS
        }
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("index", selected=request.form.get("selected", type=int)))

    get_db().execute(
        """UPDATE profile SET full_name = ?, email = ?, phone = ?, location = ?,
           postal_address = ?, linkedin_url = ?, portfolio_url = ?,
           languages = ?, hobbies = ?, certifications = ?,
           achievements = ?, resume_source = ?
           WHERE id = 1""",
        (
            request.form.get("full_name", "").strip(),
            request.form.get("email", "").strip(),
            request.form.get("phone", "").strip(),
            request.form.get("location", "").strip(),
            request.form.get("postal_address", "").strip(),
            profile_urls["linkedin_url"],
            profile_urls["portfolio_url"],
            *(request.form.get(field, "").strip() for field in PROFILE_EXTRA_FIELDS),
            resume_source,
        ),
    )
    get_db().commit()
    flash(
        "Texte du PDF extrait et enregistre sur cet ordinateur. Verifie le texte dans le profil."
        if pdf_uploaded else "Profil enregistre sur cet ordinateur.",
        "success",
    )
    return redirect(url_for("index", selected=request.form.get("selected", type=int)))


@app.post("/reference-documents")
def upload_reference_documents():
    uploads = [item for item in request.files.getlist("reference_files") if item.filename]
    if not uploads:
        flash("Choisis au moins un CV de reference a ajouter.", "error")
        return redirect(url_for("index", selected=request.form.get("selected", type=int)))

    prepared = []
    for upload in uploads:
        filename = Path(upload.filename.replace("\\", "/")).name[:200]
        try:
            content = extract_reference_document(filename, upload.read())
        except ValueError as error:
            flash(f"{filename} : {error}", "error")
            return redirect(url_for("index", selected=request.form.get("selected", type=int)))
        prepared.append((filename, content))

    database = get_db()
    inserted = 0
    duplicates = 0
    now = datetime.now().isoformat(timespec="seconds")
    for filename, content in prepared:
        duplicate = database.execute(
            "SELECT 1 FROM reference_documents WHERE content = ? LIMIT 1", (content,)
        ).fetchone()
        if duplicate:
            duplicates += 1
            continue
        database.execute(
            "INSERT INTO reference_documents (filename, content, created_at) VALUES (?, ?, ?)",
            (filename, content, now),
        )
        inserted += 1
    database.commit()
    if inserted:
        flash(f"{inserted} CV de reference ajoute(s). Leur texte sera utilise avec ton CV principal.", "success")
    elif duplicates:
        flash("Ces documents sont deja presents dans la bibliotheque de references.", "success")
    if duplicates and inserted:
        flash(f"{duplicates} document(s) identique(s) deja present(s) n'ont pas ete ajoutes.", "success")
    return redirect(url_for("index", selected=request.form.get("selected", type=int)))


@app.post("/applications/<int:application_id>/example")
def save_style_example(application_id):
    application = get_application(application_id)
    kind = request.form.get("kind")
    if application is None or kind not in {"cv", "letter"}:
        flash("Demande d'exemple invalide.", "error")
        return redirect(url_for("index"))
    content = (application["cv_draft"] if kind == "cv" else application["letter_draft"]).strip()
    if not content:
        flash("Génère d'abord un brouillon avant de le garder comme exemple.", "error")
        return redirect(url_for("index", selected=application_id))
    database = get_db()
    if database.execute(
        "SELECT 1 FROM style_examples WHERE kind = ? AND content = ?", (kind, content)
    ).fetchone():
        flash("Ce brouillon est déjà dans tes exemples.", "success")
        return redirect(url_for("index", selected=application_id))
    label = f"{application['position']} · {application['company']}"
    database.execute(
        "INSERT INTO style_examples (kind, label, content, created_at) VALUES (?, ?, ?, ?)",
        (kind, label, content, datetime.now().isoformat(timespec="seconds")),
    )
    database.commit()
    flash("Brouillon gardé comme exemple : Qwen s'en inspirera pour le style.", "success")
    return redirect(url_for("index", selected=application_id))


@app.post("/style-examples/<int:example_id>/delete")
def delete_style_example(example_id):
    database = get_db()
    database.execute("DELETE FROM style_examples WHERE id = ?", (example_id,))
    database.commit()
    flash("Exemple retiré.", "success")
    return redirect(url_for("index", selected=request.form.get("selected", type=int)))


@app.post("/reference-documents/<int:document_id>/delete")
def delete_reference_document(document_id):
    database = get_db()
    document = database.execute(
        "SELECT filename FROM reference_documents WHERE id = ?", (document_id,)
    ).fetchone()
    if document is None:
        flash("Ce document de reference n'existe pas.", "error")
        return redirect(url_for("index"))
    database.execute("DELETE FROM reference_documents WHERE id = ?", (document_id,))
    database.commit()
    flash(f"Document « {document['filename']} » retire des sources de reference.", "success")
    return redirect(url_for("index", selected=request.form.get("selected", type=int)))


def parse_match(answer):
    score = re.search(r"(?im)^\s*SCORE\s*[:=]\s*(\d{1,3})", answer)
    reason = re.search(r"(?im)^\s*RAISON\s*[:=]\s*(.+)$", answer)
    if not score:
        return None, ""
    return min(int(score.group(1)), 100), (reason.group(1).strip()[:300] if reason else "")


def score_offer(profile, offer):
    """Note la proximité entre l'offre et le profil, CV et certifications compris."""
    resume = (profile["resume_source"] or "").strip()
    certifications = profile_list(profile, "certifications")
    if not resume and not certifications:
        return None, ""
    evidence = f"CV DU CANDIDAT\n{resume[:2500] or 'Non renseigne.'}"
    if certifications:
        evidence += (
            "\n\nCERTIFICATIONS ET COMPETENCES ASSOCIEES\n"
            + "\n".join(f"- {item}" for item in certifications)
            + "\nCe sont des formations, pas des experiences professionnelles."
        )
    prompt = (
        "Tu évalues si une offre d'emploi correspond au domaine du candidat, en lisant le contenu "
        "de l'offre et pas seulement son titre. Domaine visé : " + MATCH_DOMAIN_HINT + ". "
        "Un poste de mathématiques, de statistiques ou de modélisation compte même sans le mot « data ».\n"
        "Tiens compte des certifications uniquement lorsque leurs compétences correspondent aux missions ; "
        "ne les confonds pas avec une expérience professionnelle.\n"
        "Réponds exactement sur deux lignes :\nSCORE: un entier de 0 à 100\nRAISON: une phrase courte.\n\n"
        f"{evidence}\n\n"
        f"OFFRE : {offer['position']} chez {offer['company']}\n{offer['description'][:2500]}"
    )
    try:
        return parse_match(generate_with_ollama(prompt, 120))
    except RuntimeError:
        return None, ""


def known_offer_keys():
    return {
        job_import.url_key(row["url"])
        for row in get_db().execute("SELECT url FROM applications WHERE url != ''").fetchall()
    }


def offer_from_snippet(snippet):
    """Offre minimale construite depuis la carte de l'e-mail quand la page est inaccessible."""
    lines = [line.strip() for line in snippet.splitlines() if line.strip()]
    if len(lines) < 2:
        return None
    location = next((line.strip("- ").strip() for line in lines[1:4] if re.search(r"\(\d{2,3}\)", line)), "")
    company = re.sub(r"\s+\d[.,]\d$", "", lines[1]).strip()
    return {
        "position": lines[0][:200], "company": company[:120], "location": location[:120],
        "description": "[Extrait de l'e-mail d'alerte : ouvre l'offre pour la description complète.]\n"
                       + "\n".join(lines),
    }


def import_offer_links(items):
    """Importe des offres (liens ou {'url', 'snippet'}) ; renvoie (ajoutées, doublons, échecs)."""
    profile = get_profile()
    known = known_offer_keys()
    added, duplicates, failures = [], 0, []
    for item in items:
        url, snippet = (item, "") if isinstance(item, str) else (item["url"], item.get("snippet", ""))
        if job_import.url_key(url) in known:
            duplicates += 1
            continue
        if len(added) + len(failures) >= MAX_LINKS_PER_IMPORT:
            break
        try:
            offer = job_import.import_offer(url)
        except job_import.OfferError as error:
            offer = offer_from_snippet(snippet) if snippet else None
            if offer is None:
                failures.append((url, str(error)))
                continue
        score, reason = score_offer(profile, offer)
        now = datetime.now().isoformat(timespec="seconds")
        get_db().execute(
            """INSERT INTO applications
               (company, position, location, url, job_description, match_score, match_reason,
                created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (offer["company"] or "À préciser", offer["position"], offer["location"], url,
             offer["description"], score, reason, now, now),
        )
        get_db().commit()
        known.add(job_import.url_key(url))
        added.append(offer["position"])
    return added, duplicates, failures


def flash_import_summary(added, duplicates, failures):
    if added:
        flash(f"{len(added)} offre(s) ajoutée(s) et notée(s) par Qwen.", "success")
    if duplicates:
        flash(f"{duplicates} offre(s) déjà présente(s), ignorée(s).", "success")
    for url, reason in failures[:5]:
        flash(f"Non importée ({reason}) : {url[:90]}", "error")
    if len(failures) > 5:
        flash(f"… et {len(failures) - 5} autre(s) lien(s) non importé(s).", "error")
    if not (added or duplicates or failures):
        flash("Aucun lien d'offre trouvé.", "error")


@app.post("/import-links")
def import_links():
    urls = job_import.extract_urls(request.form.get("links", ""))
    if not urls:
        flash("Colle au moins un lien http(s) d'offre.", "error")
        return redirect(url_for("index"))
    flash_import_summary(*import_offer_links(urls))
    return redirect(url_for("index"))


@app.post("/import-alerts")
def import_alerts():
    try:
        urls = job_import.fetch_alert_links()
    except job_import.OfferError as error:
        flash(str(error), "error")
        return redirect(url_for("index"))
    if not urls:
        flash("Aucune offre trouvée dans tes alertes des 7 derniers jours.", "error")
        return redirect(url_for("index"))
    flash_import_summary(*import_offer_links(urls))
    return redirect(url_for("index"))


def get_extension_token():
    database = get_db()
    row = database.execute("SELECT value FROM settings WHERE key = 'extension_token'").fetchone()
    if row:
        return row["value"]
    token = secrets.token_urlsafe(24)
    database.execute("INSERT INTO settings (key, value) VALUES ('extension_token', ?)", (token,))
    database.commit()
    return token


def score_in_background(application_id):
    def task():
        with app.app_context():
            application = get_application(application_id)
            if application is None:
                return
            score, reason = score_offer(get_profile(), {
                "position": application["position"], "company": application["company"],
                "description": application["job_description"],
            })
            if score is not None:
                get_db().execute(
                    "UPDATE applications SET match_score = ?, match_reason = ? WHERE id = ?",
                    (score, reason, application_id),
                )
                get_db().commit()

    threading.Thread(target=task, daemon=True).start()


@app.post("/api/extension/offer")
def extension_offer():
    token = request.headers.get("Authorization", "").removeprefix("Bearer ").strip()
    if not token or not hmac.compare_digest(token, get_extension_token()):
        return jsonify(error="Jeton invalide : copie-le depuis l'app."), 401
    data = request.get_json(silent=True) or {}
    url = str(data.get("url", "")).strip()[:2000]
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return jsonify(error="Lien de l'offre invalide."), 400
    description = str(data.get("description", "")).strip()[:job_import.MAX_DESCRIPTION_CHARS]
    if len(description) < 100:
        return jsonify(error="Description introuvable : sélectionne le texte de l'offre puis réessaie."), 400
    now = datetime.now().isoformat(timespec="seconds")
    key = job_import.url_key(url)
    existing = next(
        (row for row in get_db().execute("SELECT id, url, job_description FROM applications WHERE url != ''")
         if job_import.url_key(row["url"]) == key),
        None,
    )
    if existing:
        # Un clic sur l'extension corrige une description importée par erreur (e-mail d'alerte, autre page).
        if existing["job_description"].strip() == description:
            return jsonify(status="duplicate"), 200
        get_db().execute(
            "UPDATE applications SET job_description = ?, updated_at = ? WHERE id = ?",
            (description, now, existing["id"]),
        )
        get_db().commit()
        score_in_background(existing["id"])
        return jsonify(status="updated", id=existing["id"]), 200
    cursor = get_db().execute(
        """INSERT INTO applications
           (company, position, location, url, job_description, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (str(data.get("company", "")).strip()[:120] or "À préciser",
         str(data.get("title", "")).strip()[:200] or "Offre importée",
         str(data.get("location", "")).strip()[:120], url, description, now, now),
    )
    get_db().commit()
    score_in_background(cursor.lastrowid)
    return jsonify(status="added", id=cursor.lastrowid), 201


def import_alerts_at_startup():
    with app.app_context():
        try:
            import_offer_links(job_import.fetch_alert_links())
        except job_import.OfferError as error:
            app.logger.warning("Import des alertes ignoré : %s", error)


@app.post("/applications")
def add_application():
    company = request.form.get("company", "").strip()
    position = request.form.get("position", "").strip()
    offer_url = request.form.get("url", "").strip()
    if not company or not position:
        flash("Indique au minimum l'entreprise et le poste.", "error")
        return redirect(url_for("index"))
    if offer_url:
        parsed_url = urlsplit(offer_url)
        if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
            flash("Le lien de l'offre doit commencer par http:// ou https://.", "error")
            return redirect(url_for("index"))
    now = datetime.now().isoformat(timespec="seconds")
    cursor = get_db().execute(
        """INSERT INTO applications
           (company, position, location, url, company_address, job_description, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (company, position, request.form.get("location", "").strip(),
         offer_url, request.form.get("company_address", "").strip(),
         request.form.get("job_description", "").strip(), now, now),
    )
    get_db().commit()
    flash("Offre ajoutee au suivi.", "success")
    return redirect(url_for("index", selected=cursor.lastrowid))


@app.post("/applications/<int:application_id>/delete")
def delete_application(application_id):
    application = get_application(application_id)
    if application is None:
        flash("Cette candidature n'existe pas.", "error")
        return redirect(url_for("index"))
    get_db().execute("DELETE FROM applications WHERE id = ?", (application_id,))
    get_db().commit()
    flash(f"Offre « {application['position']} » supprimée.", "success")
    return redirect(url_for("index"))


@app.post("/applications/purge")
def purge_applications():
    days = request.form.get("days", type=int)
    max_score = request.form.get("max_score", type=int)
    if days is None or days < 1:
        flash("Indique un nombre de jours valide (1 ou plus).", "error")
        return redirect(url_for("index"))
    cutoff = (datetime.now() - timedelta(days=days)).isoformat(timespec="seconds")
    query = "DELETE FROM applications WHERE status = 'A analyser' AND created_at < ?"
    params = [cutoff]
    if max_score is not None:
        query += " AND (match_score IS NULL OR match_score <= ?)"
        params.append(max_score)
    deleted = get_db().execute(query, params).rowcount
    get_db().commit()
    flash(f"{deleted} offre(s) à analyser supprimée(s). Les candidatures déjà suivies sont conservées.", "success")
    return redirect(url_for("index"))


@app.post("/applications/<int:application_id>/update")
def update_application(application_id):
    application = get_application(application_id)
    if application is None:
        flash("Cette candidature n'existe pas.", "error")
        return redirect(url_for("index"))
    company = request.form.get("company")
    if company is not None:
        company = company.strip()[:120]
        if not company:
            flash("Indique le nom de l'entreprise.", "error")
            return redirect(url_for("index", selected=application_id))
    else:
        company = application["company"]
    status = request.form.get("status", "A analyser")
    if status not in STATUSES:
        status = "A analyser"
    get_db().execute(
        """UPDATE applications
           SET company = ?, status = ?, notes = ?, reply = ?, company_address = ?, updated_at = ?
           WHERE id = ?""",
        (
            company,
            status,
            request.form.get("notes", "").strip(),
            request.form.get("reply", "").strip(),
            request.form.get("company_address", "").strip(),
            datetime.now().isoformat(timespec="seconds"),
            application_id,
        ),
    )
    get_db().commit()
    flash("Suivi mis a jour.", "success")
    return redirect(url_for("index", selected=application_id))


@app.post("/applications/<int:application_id>/draft")
def generate_draft(application_id):
    application = get_application(application_id)
    profile = get_profile()
    reference_documents = get_reference_documents()
    kind = request.form.get("kind")
    if application is None or kind not in {"cv", "letter"}:
        flash("Demande de brouillon invalide.", "error")
        return redirect(url_for("index"))
    if not profile["resume_source"].strip() and not reference_documents:
        flash("Ajoute d'abord au moins un CV de reference dans le profil.", "error")
        return redirect(url_for("index", selected=application_id))
    if not application["job_description"].strip():
        flash("Ajoute la description de l'offre avant de generer un brouillon.", "error")
        return redirect(url_for("index", selected=application_id))
    if application["job_description"].lstrip().startswith("[Extrait de l'e-mail d'alerte"):
        flash(
            "La description de cette offre n'est qu'un extrait de l'e-mail d'alerte. "
            "Ouvre l'offre, colle la description complète, puis relance la génération.",
            "error",
        )
        return redirect(url_for("index", selected=application_id))
    provider = request.form.get("provider", "ollama")
    if provider not in {"ollama", "copilot"}:
        flash("Choisis un moteur de generation valide.", "error")
        return redirect(url_for("index", selected=application_id))
    if provider == "copilot" and request.form.get("allow_cloud") != "yes":
        flash(
            "Confirme l'envoi de ton CV et de l'offre à Copilot avant la génération.",
            "error",
        )
        return redirect(url_for("index", selected=application_id))
    language = resolve_language(request.form.get("language", "auto"), application)
    try:
        max_tokens = 900 if kind == "cv" else 600
        examples = () if provider == "copilot" or language == "en" else get_style_examples()
        prompt = build_prompt(kind, profile, application, reference_documents, examples, language)
        if provider == "copilot":
            draft = generate_with_copilot(prompt, max_tokens)
        elif select_style_examples(kind, examples):
            draft = generate_with_ollama(prompt, max_tokens, num_ctx=OLLAMA_CONTEXT_WITH_EXAMPLES)
        else:
            draft = generate_with_ollama(prompt, max_tokens)
        draft = clean_cv_draft(draft) if kind == "cv" else clean_letter_draft(draft)
        if not draft.strip():
            if provider == "copilot":
                raise RuntimeError(
                    "Copilot a renvoyé une réponse vide après nettoyage. Aucun brouillon n'a été enregistré. "
                    "Réessaie ou choisis Qwen local."
                )
            raise RuntimeError(
                "Le modèle a renvoyé une réponse vide après nettoyage. Aucun brouillon n'a été enregistré."
            )
        if looks_like_refusal(draft):
            raise RuntimeError(
                "Le modèle n'a pas pu rédiger : il réclame plus d'informations sur l'offre. "
                "Complète la description de l'offre puis relance la génération."
            )
    except RuntimeError as error:
        flash(str(error), "error")
        return redirect(url_for("index", selected=application_id))
    field = "cv_draft" if kind == "cv" else "letter_draft"
    get_db().execute(
        f"UPDATE applications SET {field} = ?, updated_at = ? WHERE id = ?",
        (draft, datetime.now().isoformat(timespec="seconds"), application_id),
    )
    get_db().commit()
    flash("Brouillon genere. Relis-le et verifie chaque information avant utilisation.", "success")
    return redirect(url_for("index", selected=application_id))


@app.get("/applications/<int:application_id>/cv-preview")
def cv_preview(application_id):
    application = get_application(application_id)
    if application is None or not application["cv_draft"]:
        return "Genere d'abord le CV.", 404
    layout = parse_cv_layout({key: request.args.get(key) for key in CV_LAYOUT_DEFAULTS if key in request.args}
                             or application["cv_layout"])
    try:
        pdf = compile_latex(
            make_cv_latex(get_profile(), application, application["cv_draft"], layout),
            one_page=True,
        )
    except RuntimeError as error:
        app.logger.warning("Aperçu LaTeX indisponible, PDF simple utilise: %s", error)
        pdf = make_pdf(get_profile(), application, application["cv_draft"], "cv")
    return send_file(pdf, mimetype="application/pdf")


@app.post("/applications/<int:application_id>/cv-layout")
def save_cv_layout(application_id):
    if get_application(application_id) is None:
        flash("Candidature introuvable.", "error")
        return redirect(url_for("index"))
    layout = parse_cv_layout({key: request.form.get(key) for key in CV_LAYOUT_DEFAULTS})
    get_db().execute(
        "UPDATE applications SET cv_layout = ?, updated_at = ? WHERE id = ?",
        (json.dumps(layout), datetime.now().isoformat(timespec="seconds"), application_id),
    )
    get_db().commit()
    flash("Mise en page validee : elle sera utilisee pour les telechargements PDF et LaTeX.", "success")
    return redirect(url_for("index", selected=application_id))


@app.get(
    "/applications/<int:application_id>/download/<kind>",
    defaults={"file_format": "docx"},
)
@app.get("/applications/<int:application_id>/download/<kind>/<file_format>")
def download_draft(application_id, kind, file_format):
    application = get_application(application_id)
    if (
        application is None or kind not in {"cv", "letter"}
        or file_format not in {"docx", "pdf", "tex"}
        or (file_format == "tex" and kind not in {"cv", "letter"})
    ):
        flash("Document introuvable.", "error")
        return redirect(url_for("index"))
    content = application["cv_draft" if kind == "cv" else "letter_draft"]
    if not content:
        flash("Genere d'abord ce brouillon.", "error")
        return redirect(url_for("index", selected=application_id))
    safe_name = export_file_name(get_profile(), application, kind, detect_language(content))
    profile = get_profile()
    if file_format == "tex":
        source = (
            make_cv_latex(profile, application, content, application["cv_layout"])
            if kind == "cv" else make_letter_latex(profile, application, content)
        )
        output = io.BytesIO(source.encode("utf-8"))
        mimetype = "application/x-tex"
    elif file_format == "pdf":
        output = None
        if kind == "cv":
            try:
                output = compile_latex(
                    make_cv_latex(profile, application, content, application["cv_layout"]),
                    one_page=True,
                )
            except RuntimeError as error:
                app.logger.warning("PDF LaTeX indisponible, PDF simple utilise: %s", error)
        else:
            try:
                output = compile_latex(make_letter_latex(profile, application, content))
            except RuntimeError as error:
                app.logger.warning("PDF LaTeX indisponible, PDF simple utilise: %s", error)
        if output is None:
            output = make_pdf(profile, application, content, kind)
        try:
            reader = PdfReader(output, strict=True)
            pages = reader.pages
            text = "\n".join(page.extract_text() or "" for page in pages).strip()
        except (PdfReadError, OSError, ValueError) as error:
            app.logger.exception("PDF export failed validation")
            flash("Le PDF généré est illisible. Réessaie l’export ou télécharge le DOCX.", "error")
            return redirect(url_for("index", selected=application_id))
        if not pages or not text or (kind == "cv" and len(pages) != 1):
            app.logger.error(
                "PDF export validation failed: kind=%s pages=%s text=%s",
                kind, len(pages), bool(text),
            )
            flash(
                "Le PDF généré est vide, illisible ou dépasse une page. "
                "Régénère le CV ou télécharge le DOCX.",
                "error",
            )
            return redirect(url_for("index", selected=application_id))
        output.seek(0)
        mimetype = "application/pdf"
    else:
        output = (
            make_cv_docx(profile, application, content)
            if kind == "cv"
            else make_letter_docx(profile, application, content)
        )
        mimetype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return send_file(
        output, as_attachment=True, download_name=f"{safe_name}.{file_format}",
        mimetype=mimetype,
    )


@app.get("/export.xlsx")
def export_xlsx():
    rows = get_db().execute(
        "SELECT * FROM applications ORDER BY created_at DESC, id DESC"
    ).fetchall()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Candidatures"
    headers = ["Entreprise", "Poste", "Localisation", "Lien", "Statut", "Date ajoutee",
               "Description de l'offre", "Notes", "Reponse recue", "CV adapte", "Lettre de motivation"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="21493F")
    for row in rows:
        sheet.append([row["company"], row["position"], row["location"], row["url"], row["status"],
                      row["created_at"], row["job_description"], row["notes"], row["reply"],
                      row["cv_draft"], row["letter_draft"]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 24, "B": 28, "C": 20, "D": 42, "E": 24, "F": 22}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="suivi_candidatures.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/health")
def health():
    return {"status": "ok", "ollama": ollama_status()}


with app.app_context():
    init_db()


if __name__ == "__main__":
    if os.environ.get("AUTO_IMPORT_ALERTS") == "1" and job_import.gmail_settings()["configured"]:
        threading.Thread(target=import_alerts_at_startup, daemon=True).start()
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)