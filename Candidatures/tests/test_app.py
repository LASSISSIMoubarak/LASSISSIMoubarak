import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from urllib.error import URLError
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import app as candidatures_app


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        candidatures_app.app.config.update(
            TESTING=True,
            DATABASE=str(Path(self.temp_dir.name) / "test.sqlite3"),
        )
        with candidatures_app.app.app_context():
            candidatures_app.init_db()
        self.client = candidatures_app.app.test_client()

    def post_form(self, path, data):
        self.client.get("/")
        with self.client.session_transaction() as session:
            data["csrf_token"] = session["csrf_token"]
        return self.client.post(path, data=data)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_add_offer_and_update_status(self):
        response = self.post_form(
            "/applications",
            data={
                "company": "Exemple SA",
                "position": "Analyste",
                "job_description": "Analyser des donnees.",
            },
        )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            application = candidatures_app.get_db().execute(
                "SELECT * FROM applications"
            ).fetchone()
            application_id = application["id"]
        response = self.post_form(
            f"/applications/{application_id}/update",
            data={
                "status": "Candidature envoyee",
                "notes": "Relancer lundi",
                "reply": "",
                "company": "Entreprise corrigée",
                "company_address": "20 avenue Exemple",
            },
        )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            application = candidatures_app.get_application(application_id)
            self.assertEqual(application["status"], "Candidature envoyee")
            self.assertEqual(application["company"], "Entreprise corrigée")
            self.assertEqual(application["notes"], "Relancer lundi")
            self.assertEqual(application["company_address"], "20 avenue Exemple")

    def test_application_title_removes_company_and_linkedin_suffix_for_documents(self):
        application = {
            "position": "Data Scientist /AI Engineer | Eulidia | LinkedIn",
            "company": "Eulidia",
            "location": "",
        }
        self.assertEqual(
            candidatures_app.cv_role_title(application, upper=False),
            "Data Scientist /AI Engineer",
        )

    def test_profile_and_company_postal_addresses_are_saved(self):
        response = self.post_form(
            "/profile",
            data={
                "full_name": "Nom Exemple",
                "postal_address": "12 rue des Donnees\n49000 Angers",
                "resume_source": "SQL et Python",
            },
        )
        self.assertEqual(response.status_code, 302)
        response = self.post_form(
            "/applications",
            data={
                "company": "Exemple",
                "position": "Data Analyst",
                "company_address": "20 avenue des Donnees\n49000 Angers",
            },
        )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            profile = candidatures_app.get_profile()
            application = candidatures_app.get_db().execute(
                "SELECT * FROM applications"
            ).fetchone()
            self.assertIn("12 rue des Donnees", profile["postal_address"])
            self.assertIn("20 avenue des Donnees", application["company_address"])

    def test_profile_social_links_are_saved_and_rendered_in_cv_header(self):
        response = self.post_form(
            "/profile",
            data={
                "linkedin_url": "linkedin.com/in/moubarak",
                "portfolio_url": "https://portfolio.example/projects",
            },
        )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            profile = candidatures_app.get_profile()
            self.assertEqual(profile["linkedin_url"], "https://linkedin.com/in/moubarak")
            self.assertEqual(profile["portfolio_url"], "https://portfolio.example/projects")
            source = candidatures_app.make_cv_latex(
                profile, {"position": "Data", "company": "X"}, "PROFIL\nProfil court."
            )
            self.assertIn(r"\href{https://linkedin.com/in/moubarak}{LinkedIn}", source)
            self.assertIn(r"\href{https://portfolio.example/projects}{Portfolio}", source)
            if candidatures_app.shutil.which("xelatex"):
                pdf_text = PdfReader(candidatures_app.compile_latex(source)).pages[0].extract_text()
                self.assertIn("LinkedIn", pdf_text)
                self.assertIn("Portfolio", pdf_text)

    def test_profile_rejects_invalid_social_link(self):
        response = self.post_form("/profile", data={"linkedin_url": "javascript:alert(1)"})
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            self.assertEqual(candidatures_app.get_profile()["linkedin_url"], "")

    def test_export_contains_headers_and_offer(self):
        self.post_form(
            "/applications",
            data={"company": "Exemple SA", "position": "Analyste"},
        )
        response = self.client.get("/export.xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        sheet = workbook["Candidatures"]
        self.assertEqual(sheet.cell(1, 1).value, "Entreprise")
        self.assertEqual(sheet.cell(2, 1).value, "Exemple SA")

    def test_home_page_renders_empty_state(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Le tableau est prêt".encode(), response.data)

    def test_project_library_and_detail_render_seeded_projects(self):
        library = self.client.get("/projects")
        self.assertEqual(library.status_code, 200)
        self.assertIn("SolarWatch".encode(), library.data)
        self.assertIn("Tableau de bord Power BI".encode(), library.data)
        self.assertIn("Ajouter un autre projet".encode(), library.data)

        detail = self.client.get("/projects/solarwatch")
        self.assertEqual(detail.status_code, 200)
        self.assertIn("Projet personnel".encode(), detail.data)
        self.assertIn("pas de pannes labellisées".encode(), detail.data)
        self.assertIn("pas encore accessible à un recruteur".encode(), detail.data)

        bi_detail = self.client.get("/projects/power-bi-transactions-india")
        self.assertEqual(bi_detail.status_code, 200)
        self.assertIn("Power Automate".encode(), bi_detail.data)
        self.assertIn("transactions par secteur".encode(), bi_detail.data)

    def test_projects_can_be_added_and_updated_with_validated_public_url(self):
        with candidatures_app.app.app_context():
            project_id = candidatures_app.get_db().execute(
                "SELECT id FROM projects WHERE slug = 'solarwatch'"
            ).fetchone()["id"]

        response = self.post_form("/projects/save", data={
            "project_id": str(project_id),
            "title": "SolarWatch actualisé",
            "summary": "Résumé vérifié.",
            "description": "Périmètre documenté.",
            "skills": "Python, LSTM",
            "keywords": "solaire, prévision",
            "highlights": "Résultat réellement observé",
            "public_url": "github.com/exemple/solarwatch",
        })
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            project = candidatures_app.get_db().execute(
                "SELECT * FROM projects WHERE id = ?", (project_id,)
            ).fetchone()
            self.assertEqual(project["title"], "SolarWatch actualisé")
            self.assertEqual(project["public_url"], "https://github.com/exemple/solarwatch")
            self.assertEqual(project["highlights"], "Résultat réellement observé")

        response = self.post_form("/projects/save", data={
            "title": "Projet ajouté",
            "summary": "Résumé du projet.",
            "description": "Description du projet.",
            "skills": "Python",
            "keywords": "python, analyse",
            "highlights": "Test documenté",
            "public_url": "javascript:alert(1)",
        })
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            self.assertIsNone(candidatures_app.get_db().execute(
                "SELECT id FROM projects WHERE slug = 'projet-ajoute'"
            ).fetchone())

        response = self.post_form("/projects/save", data={
            "title": "Projet ajouté",
            "summary": "Résumé du projet.",
            "description": "Description du projet.",
            "skills": "Python",
            "keywords": "python, analyse",
            "highlights": "Test documenté",
            "public_url": "",
        })
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            self.assertIsNotNone(candidatures_app.get_db().execute(
                "SELECT id FROM projects WHERE slug = 'projet-ajoute'"
            ).fetchone())

    def test_relevant_projects_are_persisted_and_used_in_cv_prompt_and_content(self):
        response = self.post_form("/applications", data={
            "company": "TotalEnergies",
            "position": "Ingénieur Data Solaire et IA",
            "job_description": (
                "Analyser la production photovoltaïque, détecter les anomalies "
                "des onduleurs et prévoir la production avec un modèle LSTM."
            ),
        })
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            application = candidatures_app.get_db().execute(
                "SELECT * FROM applications"
            ).fetchone()
            application_id = application["id"]
            choices = candidatures_app.ensure_application_project_choices(application)
            self.assertEqual(choices[0]["slug"], "solarwatch")
            solarwatch = next(item for item in choices if item["slug"] == "solarwatch")
            self.assertEqual(solarwatch["is_selected"], 1)
            projects = candidatures_app.selected_application_projects(application_id)
            prompt = candidatures_app.build_prompt(
                "cv",
                candidatures_app.get_profile(),
                application,
                selected_projects=projects,
            )
            self.assertIn("PROJETS PERSONNELS SÉLECTIONNÉS", prompt)
            self.assertIn("Le jeu ne contient pas de pannes labellisées", prompt)

            cv = candidatures_app.inject_selected_projects("PROFIL\nData analyst.", projects)
            self.assertIn("PROJETS", cv)
            self.assertIn("Projet personnel", cv)
            self.assertIn("alertes candidates", cv)
            self.assertNotIn("PROJETS\nSolarWatch\nUn résultat inventé", cv)
            self.assertLessEqual(len(projects), 2)

    def test_cv_keeps_original_single_project_layout_for_two_selected_projects(self):
        projects = [
            {
                "title": "Projet le plus pertinent",
                "summary": "Surveillance photovoltaïque.",
                "skills": "Python, pandas, LSTM",
                "highlights": "Résultat mesuré sur test chronologique",
                "public_url": "https://github.com/example/solar",
            },
            {
                "title": "Projet secondaire",
                "summary": "Autre projet de machine learning.",
                "skills": "Python, scikit-learn",
                "highlights": "Résultat secondaire",
                "public_url": "",
            },
        ]
        content = candidatures_app.inject_selected_projects(
            "PROFIL\nData analyst.\nFORMATION\nMaster Data", projects
        )
        self.assertIn("Projet le plus pertinent", content)
        self.assertIn("Résultat mesuré sur test chronologique", content)
        self.assertIn("https://github.com/example/solar", content)
        self.assertNotIn("Projet secondaire", content)

        source = candidatures_app.make_cv_latex(
            {"full_name": "Nom Exemple", "email": "", "phone": "", "location": ""},
            {"position": "Data Analyst", "company": "Exemple"},
            content,
        )
        self.assertIn(r"\cvsection{Projets}", source)
        self.assertIn(r"\href{\detokenize{https://github.com/example/solar}}", source)
        self.assertNotIn("Projet secondaire", source)
        if candidatures_app.shutil.which("xelatex"):
            pdf = PdfReader(candidatures_app.compile_latex(source, one_page=True))
            self.assertEqual(len(pdf.pages), 1)
            self.assertIn("Projet le plus pertinent", pdf.pages[0].extract_text())

    def test_cv_project_title_uses_a_single_plain_hyphen(self):
        project = {
            "title": "Pipeline MLOps — classification du cancer du sein",
            "summary": "Pipeline reproductible.",
            "skills": "Python, scikit-learn",
            "highlights": "",
            "public_url": "",
        }
        cv = candidatures_app.inject_selected_projects("PROFIL\nData analyst.", [project])
        self.assertIn("Pipeline MLOps - classification du cancer du sein", cv)
        self.assertNotIn("—", cv)
        self.assertNotIn("–", cv)

    def test_cv_injection_replaces_all_duplicate_project_sections(self):
        draft = (
            "PROFIL\nData analyst.\n"
            "PROJETS\nAncien projet un\n- Description une.\n"
            "FORMATION\nMaster Data\n"
            "PROJECTS\nAncien projet deux\n- Description deux.\n"
        )
        projects = [{
            "title": "Projet sélectionné",
            "summary": "Résumé vérifié.",
            "skills": "Python",
            "highlights": "Résultat vérifiable",
            "public_url": "",
        }]
        updated = candidatures_app.inject_selected_projects(draft, projects)
        self.assertEqual(
            sum(
                candidatures_app.CV_SECTION_ALIASES.get(
                    candidatures_app.normalize_heading(line),
                    candidatures_app.normalize_heading(line),
                ) == "PROJETS"
                for line in updated.splitlines()
            ),
            1,
        )
        self.assertIn("Projet sélectionné", updated)
        self.assertNotIn("Ancien projet un", updated)
        self.assertNotIn("Ancien projet deux", updated)
        self.assertLess(updated.index("PROJETS"), updated.index("FORMATION"))

    def test_application_project_selection_is_limited_to_two(self):
        self.post_form("/applications", data={
            "company": "Exemple",
            "position": "Data",
            "job_description": "machine learning data science Python analyse",
        })
        with candidatures_app.app.app_context():
            application_id = candidatures_app.get_db().execute(
                "SELECT id FROM applications"
            ).fetchone()["id"]
            choices = candidatures_app.ensure_application_project_choices(
                candidatures_app.get_application(application_id)
            )
            project_ids = [str(choice["id"]) for choice in choices[:3]]
        response = self.post_form(
            f"/applications/{application_id}/projects",
            data={"project_ids": project_ids},
        )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            selected = candidatures_app.selected_application_projects(application_id)
            self.assertLessEqual(len(selected), 2)

    def test_project_primary_choice_is_respected_and_generation_saves_unsaved_selection(self):
        self.post_form("/profile", data={"resume_source": "Data analyst, Python, machine learning."})
        self.post_form("/applications", data={
            "company": "Exemple",
            "position": "Ingénieur data énergie",
            "job_description": (
                "Projet solaire photovoltaïque, détection d'anomalies, prévision LSTM. "
                "Pipeline reproductible machine learning avec scikit-learn et régression logistique."
            ),
        })
        with candidatures_app.app.app_context():
            application = candidatures_app.get_db().execute(
                "SELECT * FROM applications"
            ).fetchone()
            application_id = application["id"]
            choices = candidatures_app.ensure_application_project_choices(application)
            solarwatch = next(item for item in choices if item["slug"] == "solarwatch")
            mlops = next(item for item in choices if item["slug"] == "mlops-breast-cancer")
            error = candidatures_app.save_application_project_selection(
                application,
                [str(solarwatch["id"]), str(mlops["id"])],
                str(solarwatch["id"]),
            )
            self.assertIsNone(error)
            selected = candidatures_app.selected_application_projects(application_id)
            self.assertEqual(selected[0]["slug"], "solarwatch")

        page = self.client.get(f"/?selected={application_id}")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'name="primary_project_id"', page.data)
        self.assertIn(b"project_selection_submitted", page.data)

        desired_id = str(mlops["id"])
        other_selected_id = str(solarwatch["id"])
        with patch(
            "app.generate_with_ollama",
            return_value="PROFIL\nAnalyste de données.\nFORMATION\nMaster Data.",
        ):
            response = self.post_form(
                f"/applications/{application_id}/draft",
                data={
                    "kind": "cv",
                    "provider": "ollama",
                    "project_selection_submitted": "1",
                    "project_ids": [other_selected_id, desired_id],
                    "primary_project_id": desired_id,
                },
            )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            application = candidatures_app.get_application(application_id)
            selected = candidatures_app.selected_application_projects(application_id)
            self.assertEqual(
                [project["slug"] for project in selected],
                ["mlops-breast-cancer", "solarwatch"],
            )
            self.assertIn("Pipeline MLOps - classification du cancer du sein", application["cv_draft"])
            self.assertNotIn("SolarWatch - surveillance", application["cv_draft"])

    def test_home_page_renders_one_application_table_header(self):
        self.post_form(
            "/applications",
            data={"company": "Exemple SA", "position": "Analyste"},
        )
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data.count(b"<th>Poste</th>"), 1)

    def test_post_without_csrf_token_is_rejected(self):
        response = self.client.post(
            "/applications", data={"company": "X", "position": "Y"}
        )
        self.assertEqual(response.status_code, 400)

    def test_pdf_resume_text_is_extracted_and_saved(self):
        page = Mock()
        page.extract_text.return_value = "Nom Exemple\nCompetence Python"
        reader = Mock()
        reader.pages = [page]
        with patch("app.PdfReader", return_value=reader):
            response = self.post_form(
                "/profile",
                data={
                    "full_name": "Nom Exemple",
                    "resume_pdf": (io.BytesIO(b"%PDF-test"), "cv.pdf"),
                },
            )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            profile = candidatures_app.get_profile()
            self.assertIn("Competence Python", profile["resume_source"])

    def test_latex_reference_extracts_readable_cv_content(self):
        source = r"""\documentclass{article}
\newcommand{\cvsection}[1]{command definition}
\begin{document}
\textcolor{accent}{DATA SCIENTIST}
\cvsection{Expériences professionnelles}
\cventry{DATA SCIENTIST}{ACME}{2024 - 2025}
\begin{cvitems}
\item Analyser les données avec Python.
\item Utiliser SQL et Power BI.
\end{cvitems}
\cvsection{Compétences clés}
\begin{cvitems}
\item Modélisation statistique
\end{cvitems}
\end{document}"""
        extracted = candidatures_app.extract_latex_text(source)
        self.assertIn("Expériences professionnelles", extracted)
        self.assertIn("DATA SCIENTIST", extracted)
        self.assertIn("Analyser les données avec Python.", extracted)
        self.assertIn("Modélisation statistique", extracted)
        self.assertNotIn("documentclass", extracted)
        self.assertNotIn("command definition", extracted)
        self.assertNotIn("accent", extracted)

    def test_reference_documents_are_combined_deduplicated_and_removable(self):
        response = self.post_form(
            "/reference-documents",
            data={
                "reference_files": [
                    (
                        io.BytesIO(
                            b"\\begin{document}\\cvsection{Competences}\\item SQL\\item Python\\end{document}"
                        ),
                        "cv-data.tex",
                    ),
                    (io.BytesIO(b"Python\nPower BI\nR"), "cv-bi.txt"),
                ]
            },
        )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            documents = candidatures_app.get_reference_documents()
            self.assertEqual(len(documents), 2)
            profile = dict(candidatures_app.get_profile())
            profile["resume_source"] = "SQL\nPython"
            prompt = candidatures_app.build_prompt(
                "cv",
                profile,
                {"position": "Data Analyst", "company": "Exemple", "location": "", "job_description": "SQL"},
                documents,
            )
            self.assertIn("Power BI", prompt)
            self.assertEqual(prompt.count("Python"), 1)
            document_id = documents[0]["id"]

        response = self.post_form(f"/reference-documents/{document_id}/delete", data={})
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            self.assertEqual(len(candidatures_app.get_reference_documents()), 1)

    def test_generation_can_use_reference_documents_without_primary_cv_text(self):
        self.post_form(
            "/applications",
            data={
                "company": "Exemple",
                "position": "Data Analyst",
                "job_description": "Power BI et SQL.",
            },
        )
        with candidatures_app.app.app_context():
            application_id = candidatures_app.get_db().execute(
                "SELECT id FROM applications"
            ).fetchone()["id"]
        self.post_form(
            "/reference-documents",
            data={"reference_files": [(io.BytesIO(b"SQL\nPower BI\nReporting"), "cv.txt")]},
        )
        with patch("app.generate_with_ollama", return_value="PROFIL\nAnalyste.") as generate:
            response = self.post_form(
                f"/applications/{application_id}/draft", data={"kind": "cv"}
            )
        self.assertEqual(response.status_code, 302)
        self.assertIn("Power BI", generate.call_args.args[0])
        with candidatures_app.app.app_context():
            self.assertEqual(
                candidatures_app.get_application(application_id)["cv_draft"],
                    "PROFIL\nAnalyste.\n\nPROJETS\n"
                    "Tableau de bord Power BI - transactions en Inde\n"
                    "Projet personnel\n"
                    "- Rapport interactif orienté métier pour explorer les montants "
                    "de transactions par secteur, genre, État, année et âge. Résultat "
                    "documenté : Analyse interactive des montants par secteur "
                    "d'activité et par genre\n"
                    "- Technologies : Power BI, conception de tableaux de bord, "
                    "visualisation de données, analyse métier, Power Automate, "
                    "automatisation de l'actualisation",
                )

    def test_copilot_generation_requires_consent_and_uses_selected_provider(self):
        self.post_form(
            "/profile",
            data={"resume_source": "Analyste de données, Python et SQL."},
        )
        self.post_form(
            "/applications",
            data={
                "company": "Exemple",
                "position": "Data Analyst",
                "job_description": "Python et SQL.",
            },
        )
        with candidatures_app.app.app_context():
            application_id = candidatures_app.get_db().execute(
                "SELECT id FROM applications"
            ).fetchone()["id"]

        with patch("app.generate_with_copilot", return_value="PROFIL\nAnalyste.") as generate:
            response = self.post_form(
                f"/applications/{application_id}/draft",
                data={"kind": "cv", "provider": "copilot"},
            )
            self.assertEqual(response.status_code, 302)
            generate.assert_not_called()

            response = self.post_form(
                f"/applications/{application_id}/draft",
                data={
                    "kind": "cv",
                    "provider": "copilot",
                    "allow_cloud": "yes",
                },
            )
        self.assertEqual(response.status_code, 302)
        generate.assert_called_once()
        with candidatures_app.app.app_context():
            self.assertEqual(
                candidatures_app.get_application(application_id)["cv_draft"],
                "PROFIL\nAnalyste.",
            )

    def test_empty_copilot_draft_is_not_saved(self):
        self.post_form("/profile", data={"resume_source": "SQL et Python."})
        self.post_form("/applications", data={
            "company": "Exemple", "position": "Data Analyst",
            "job_description": "SQL et Python.",
        })
        with candidatures_app.app.app_context():
            application_id = candidatures_app.get_db().execute(
                "SELECT id FROM applications"
            ).fetchone()["id"]
        with patch("app.generate_with_copilot", return_value=""):
            response = self.post_form(
                f"/applications/{application_id}/draft",
                data={"kind": "cv", "provider": "copilot", "allow_cloud": "yes"},
            )
        self.assertEqual(response.status_code, 302)
        with candidatures_app.app.app_context():
            self.assertEqual(
                candidatures_app.get_application(application_id)["cv_draft"], ""
            )
        page = self.client.get(f"/?selected={application_id}")
        self.assertIn(b"Copilot a renvoy", page.data)
        self.assertIn(b"Aucun brouillon", page.data)

    def test_style_examples_guide_qwen_only(self):
        self.post_form("/profile", data={"resume_source": "Analyste, Python."})
        self.post_form("/applications", data={
            "company": "Exemple", "position": "Data Analyst", "job_description": "Python."})
        with candidatures_app.app.app_context():
            application_id = candidatures_app.get_db().execute(
                "SELECT id FROM applications").fetchone()["id"]
        good = "PROFIL\nAnalyste rigoureux.\n" + "ligne\n" * 600
        with patch("app.generate_with_copilot", return_value=good):
            self.post_form(f"/applications/{application_id}/draft", data={
                "kind": "cv", "provider": "copilot", "allow_cloud": "yes"})
        self.post_form(f"/applications/{application_id}/example", data={"kind": "cv"})
        self.post_form(f"/applications/{application_id}/example", data={"kind": "cv"})
        with patch(
            "app.generate_with_copilot",
            return_value="Madame, Monsieur,\nVotre offre m'intéresse.\nJe vous prie d'agréer.",
        ):
            self.post_form(f"/applications/{application_id}/draft", data={
                "kind": "letter", "provider": "copilot", "allow_cloud": "yes"})
        self.post_form(f"/applications/{application_id}/example", data={"kind": "letter"})
        with candidatures_app.app.app_context():
            examples = candidatures_app.get_style_examples()
            self.assertEqual(len(examples), 2)
            application = candidatures_app.get_application(application_id)
            prompt = candidatures_app.build_prompt(
                "cv", candidatures_app.get_profile(), application, (), examples)
            self.assertIn("EXEMPLES DE CV VALIDÉS", prompt)
            self.assertIn("Analyste rigoureux", prompt)
            self.assertNotIn("Votre offre m'intéresse", prompt)
            self.assertLess(len(prompt) - len(candidatures_app.build_prompt(
                "cv", candidatures_app.get_profile(), application)), 2500)
            letter_prompt = candidatures_app.build_prompt(
                "letter", candidatures_app.get_profile(), application, (), examples)
            self.assertIn("EXEMPLES DE LETTRES DE MOTIVATION VALIDÉS", letter_prompt)
            self.assertIn("Votre offre m'intéresse", letter_prompt)
            self.assertNotIn("EXEMPLES DE CV VALIDÉS", letter_prompt)
            english_letter_prompt = candidatures_app.build_prompt(
                "letter", candidatures_app.get_profile(), application, (), examples, "en")
            self.assertIn("Votre offre m'intéresse", english_letter_prompt)

        with patch("app.generate_with_ollama", return_value="PROFIL\nOk.") as ollama:
            self.post_form(f"/applications/{application_id}/draft", data={"kind": "cv"})
        self.assertIn("Analyste rigoureux", ollama.call_args.args[0])
        self.assertEqual(ollama.call_args.kwargs["num_ctx"], candidatures_app.OLLAMA_CONTEXT_WITH_EXAMPLES)

        with patch("app.generate_with_ollama", return_value="Madame, Monsieur.\nCordialement.") as ollama:
            self.post_form(f"/applications/{application_id}/draft", data={
                "kind": "letter", "language": "en"})
        self.assertIn("Votre offre m'intéresse", ollama.call_args.args[0])
        self.assertEqual(ollama.call_args.kwargs["num_ctx"], candidatures_app.OLLAMA_CONTEXT_WITH_EXAMPLES)

        with patch("app.generate_with_copilot", return_value="Madame, Monsieur.\nCordialement.") as copilot:
            self.post_form(f"/applications/{application_id}/draft", data={
                "kind": "letter", "provider": "copilot", "allow_cloud": "yes"})
        self.assertNotIn("EXEMPLES", copilot.call_args.args[0])

        for example in examples:
            self.post_form(f"/style-examples/{example['id']}/delete", data={})
        with candidatures_app.app.app_context():
            self.assertEqual(candidatures_app.get_style_examples(), [])

    def test_copilot_generator_uses_sdk_without_local_tools(self):
        from copilot.generated.session_events import AssistantMessageData

        calls = {}

        class FakeSession:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return None

            async def send_and_wait(self, prompt, timeout):
                calls["prompt"] = prompt
                calls["timeout"] = timeout
                return SimpleNamespace(
                    data=AssistantMessageData(content="Brouillon Copilot", message_id="test")
                )

        class FakeClient:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return None

            async def create_session(self, **kwargs):
                calls["session_options"] = kwargs
                return FakeSession()

        with patch("copilot.CopilotClient", return_value=FakeClient()):
            result = candidatures_app.generate_with_copilot("Prompt de test", 100)

        self.assertEqual(result, "Brouillon Copilot")
        self.assertEqual(calls["prompt"], "Prompt de test")
        self.assertEqual(calls["timeout"], candidatures_app.COPILOT_TIMEOUT)
        self.assertEqual(calls["session_options"]["available_tools"], [])
        self.assertFalse(calls["session_options"]["enable_config_discovery"])
        self.assertEqual(
            calls["session_options"]["on_permission_request"](None, None).kind,
            "reject",
        )

    def test_ollama_retries_after_transient_connection_error(self):
        response = Mock()
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=None)
        response.read.return_value = json.dumps(
            {"message": {"content": "OK"}}
        ).encode("utf-8")
        with patch(
            "app.urlopen",
            side_effect=[URLError(ConnectionRefusedError("refused")), response],
        ) as request:
            with patch("app.time.sleep"):
                result = candidatures_app.generate_with_ollama("test", max_tokens=20)
        self.assertEqual(result, "OK")
        self.assertEqual(request.call_count, 2)

    def test_ollama_timeout_does_not_repeat_long_generation(self):
        with patch("app.urlopen", side_effect=URLError(TimeoutError("timed out"))) as request:
            with self.assertRaisesRegex(RuntimeError, "Qwen n'a pas termine"):
                candidatures_app.generate_with_ollama("test", max_tokens=20)
        self.assertEqual(request.call_count, 1)

    def test_latex_export_rejects_unreadable_pdf_output(self):
        def write_invalid_pdf(*_args, **kwargs):
            (Path(kwargs["cwd"]) / "cv.pdf").write_bytes(b"not a PDF")
            return Mock()

        with patch("app.shutil.which", return_value="xelatex"), patch(
            "app.subprocess.run", side_effect=write_invalid_pdf
        ):
            with self.assertRaisesRegex(RuntimeError, "PDF compilé est illisible"):
                candidatures_app.compile_latex(r"\documentclass{article}\begin{document}CV\end{document}")

    def test_data_cv_docx_uses_candidate_and_role_not_board_name(self):
        profile = {
            "full_name": "Nom Exemple",
            "location": "Angers",
            "email": "test@example.com",
            "phone": "0600000000",
        }
        application = {"position": "Data Analyst"}
        source = "CV adapte - Hellowork\n**PROFIL**\nAnalyste de donnees.\nCOMPETENCES CLES\n- SQL"
        document = Document(
            io.BytesIO(candidatures_app.make_cv_docx(profile, application, source).getvalue())
        )
        visible_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Nom Exemple", visible_text)
        self.assertIn("DATA ANALYST", visible_text)
        self.assertIn("COMPETENCES CLES", visible_text)
        self.assertNotIn("CV adapte - Hellowork", visible_text)

    def test_cv_draft_removes_contact_details_added_by_document_layout(self):
        draft = (
            "PROFIL\nNom : Nom Exemple\nE-mail : nom@example.com\n"
            "Telephone : 0600000000\nLocalisation : Angers\n"
            "Analyste de donnees.\nCOMPETENCES CLES\n- SQL"
        )
        cleaned = candidatures_app.clean_cv_draft(draft)
        self.assertNotIn("Nom :", cleaned)
        self.assertNotIn("E-mail :", cleaned)
        self.assertNotIn("Localisation :", cleaned)
        self.assertIn("Analyste de donnees.", cleaned)

    def test_letter_prompt_demands_offer_specific_evidence_and_short_body(self):
        profile = {
            "full_name": "Nom Exemple", "email": "nom@example.com", "phone": "",
            "location": "Angers", "resume_source": "SQL et Power BI chez ACME.",
            "certifications": "Power BI Bootcamp : DAX, visualisation interactive\n"
                              "Working with SSH : administration Linux",
        }
        application = {
            "position": "Data Analyst", "company": "Exemple", "location": "",
            "job_description": "SQL, Power BI et suivi d'indicateurs.",
        }
        prompt = candidatures_app.build_prompt("letter", profile, application)
        self.assertIn("C'est avec grand plaisir que je vous adresse ma candidature", prompt)
        self.assertIn("Ne commence jamais par « J'ai le plaisir de vous proposer »", prompt)
        self.assertIn("Power BI Bootcamp : DAX, visualisation interactive", prompt)
        self.assertIn("Working with SSH : administration Linux", prompt)
        self.assertIn("uniquement les certifications dont les competences correspondent", prompt)
        self.assertIn("sans les presenter comme une experience professionnelle", prompt)
        self.assertIn("contribution concrete attendue", prompt)
        self.assertIn("2 ou 3 exigences", prompt)
        self.assertIn("ne repete ni exemple, ni phrase, ni paragraphe", prompt)
        self.assertIn("Ne calcule jamais un nombre d'annees", prompt)
        self.assertNotIn("Toujours avec enthousiasme", prompt)
        self.assertIn("N'invente aucun fait", prompt)

    def test_letter_uses_primary_selected_project_as_offer_evidence(self):
        self.post_form(
            "/profile",
            data={"resume_source": "Data analyst, Python, statistiques."},
        )
        self.post_form(
            "/applications",
            data={
                "company": "TotalEnergies",
                "position": "Ingénieur Data Solaire",
                "job_description": "Analyser la production photovoltaïque et prévoir l'énergie solaire.",
            },
        )
        with candidatures_app.app.app_context():
            application = candidatures_app.get_db().execute(
                "SELECT * FROM applications"
            ).fetchone()
            application_id = application["id"]
            projects = {
                row["slug"]: row
                for row in candidatures_app.get_projects()
                if row["slug"] in {"solarwatch", "mlops-breast-cancer"}
            }
            error = candidatures_app.save_application_project_selection(
                application,
                [str(projects["solarwatch"]["id"]), str(projects["mlops-breast-cancer"]["id"])],
                str(projects["solarwatch"]["id"]),
            )
            self.assertIsNone(error)

        generated_letter = (
            "Je candidate au poste proposé chez TotalEnergies.\n\n"
            "Mon parcours comprend des expériences de data et de statistiques.\n\n"
            "Ces compétences répondent aux besoins du poste et me permettront de contribuer aux analyses.\n\n"
            "Je souhaite rejoindre vos équipes et contribuer à vos missions."
        )
        with patch("app.generate_with_ollama", return_value=generated_letter) as generate:
            response = self.post_form(
                f"/applications/{application_id}/draft",
                data={"kind": "letter", "provider": "ollama", "language": "fr"},
            )
        self.assertEqual(response.status_code, 302)
        prompt = generate.call_args.args[0]
        self.assertIn("Projet personnel principal choisi", prompt)
        self.assertIn("SolarWatch — surveillance de performance photovoltaïque", prompt)
        self.assertNotIn("Pipeline MLOps", prompt)
        with candidatures_app.app.app_context():
            letter = candidatures_app.get_application(application_id)["letter_draft"]
            self.assertIn(
                "SolarWatch - surveillance de performance photovoltaïque", letter
            )
            self.assertIn("en mettant en pratique Python, pandas, scikit-learn, PyTorch", letter)
            self.assertNotIn("Pipeline MLOps", letter)
            self.assertLess(
                letter.index("SolarWatch"),
                letter.index("Ces compétences répondent"),
            )

    def test_offer_matching_uses_certifications_without_misrepresenting_them(self):
        profile = {
            "resume_source": "",
            "certifications": "AWS Certified Machine Learning - Specialty : machine learning",
        }
        offer = {
            "position": "Machine Learning Engineer",
            "company": "Exemple",
            "description": "Développer et évaluer des modèles de machine learning.",
        }
        with patch("app.generate_with_ollama", return_value="SCORE: 81\nRAISON: Machine learning.") as generate:
            score, reason = candidatures_app.score_offer(profile, offer)
        self.assertEqual(score, 81)
        self.assertEqual(reason, "Machine learning.")
        prompt = generate.call_args.args[0]
        self.assertIn("AWS Certified Machine Learning - Specialty", prompt)
        self.assertIn("ne les confonds pas avec une expérience professionnelle", prompt)

    def test_letter_generation_uses_enough_tokens_for_complete_body(self):
        self.post_form(
            "/applications",
            data={"company": "Exemple", "position": "Data Analyst", "job_description": "SQL"},
        )
        with candidatures_app.app.app_context():
            application_id = candidatures_app.get_db().execute(
                "SELECT id FROM applications"
            ).fetchone()["id"]
            candidatures_app.get_db().execute(
                "UPDATE profile SET resume_source = 'SQL chez ACME' WHERE id = 1"
            )
            candidatures_app.get_db().commit()
        with patch("app.generate_with_ollama", return_value="Paragraphe de test.") as generate:
            response = self.post_form(
                f"/applications/{application_id}/draft", data={"kind": "letter"}
            )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(generate.call_args.args[1], 600)

    def test_letter_pdf_has_french_letter_layout_without_cv_headings(self):
        self.post_form(
            "/applications",
            data={
                "company": "Hellowork",
                "position": "Data Scientist",
                "job_description": "Prevision de la demande et modelisation.",
            },
        )
        with candidatures_app.app.app_context():
            database = candidatures_app.get_db()
            database.execute(
                "UPDATE profile SET full_name = ?, location = ?, email = ?, postal_address = ? WHERE id = 1",
                ("Nom Exemple", "Rennes", "nom@example.com", "12 rue des Donnees\n49000 Angers"),
            )
            application = database.execute("SELECT id FROM applications").fetchone()
            application_id = application["id"]
            database.execute(
                "UPDATE applications SET company = ?, letter_draft = ?, company_address = ? WHERE id = ?",
                (
                    "Baker Tilly",
                    "Motivation Letter\n\nMadame/Monsieur,\n\nÀ PROPOS\n"
                    "J'ai analyse des donnees et construit des modeles.\n\n"
                    "Je vous prie d’agréer, Madame, Monsieur, mes salutations distinguées.",
                    "20 avenue Exemple\n49000 Angers",
                    application_id,
                ),
            )
            database.commit()

        response = self.client.get(f"/applications/{application_id}/download/letter/pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/pdf")
        pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(response.data)).pages
        )
        self.assertIn("Nom Exemple", pdf_text)
        self.assertIn("Objet : Candidature au poste de Data Scientist", pdf_text)
        self.assertIn("Madame, Monsieur", pdf_text)
        self.assertIn("12 rue des Donnees", pdf_text)
        self.assertIn("Baker Tilly", pdf_text)
        self.assertIn("20 avenue Exemple", pdf_text)
        self.assertIn("salutations distinguées", pdf_text)
        self.assertGreaterEqual(pdf_text.count("Nom Exemple"), 2)
        self.assertNotIn("Motivation Letter", pdf_text)
        self.assertNotIn("A PROPOS", pdf_text)

    def test_cv_latex_escapes_special_characters_and_compiles(self):
        profile = {"full_name": "Jean & Co", "email": "a@b.fr", "phone": "0102", "location": "Lyon"}
        application = {"position": "Data_Analyst", "company": "X"}
        draft = (
            "PROFIL\nRigueur 100% \\input{secret} {x}\n"
            "EXPERIENCES PROFESSIONNELLES\nDATA ANALYST\nACME\n04/2026 - 09/2026\n- Dashboards #1\n"
        )
        source = candidatures_app.make_cv_latex(profile, application, draft)
        self.assertIn(r"Jean \& Co", source)
        self.assertIn(r"100\% \textbackslash{}input\{secret\} \{x\}", source)
        self.assertIn(r"\cventry{DATA ANALYST}{ACME}{04/2026 - 09/2026}", source)
        if candidatures_app.shutil.which("xelatex"):
            pdf_text = PdfReader(candidatures_app.compile_latex(source)).pages[0].extract_text()
            self.assertIn("Dashboards", pdf_text)

    def test_cv_prompt_and_layout_limit_experience_to_two_activities(self):
        profile = {
            "full_name": "Nom", "email": "", "phone": "", "location": "",
            "resume_source": "SQL, Python, Power BI",
        }
        application = {
            "position": "Data Analyst", "company": "Exemple",
            "location": "", "job_description": "SQL et Power BI.",
        }
        prompt = candidatures_app.build_prompt("cv", profile, application)
        self.assertIn("au plus 2 activites", prompt)
        self.assertIn("au plus les 3 experiences les plus pertinentes", prompt)
        self.assertIn("hautement significatives", prompt)
        draft = (
            "EXPERIENCES PROFESSIONNELLES\nDATA ANALYST\nACME\n2024 – 2025\n"
            "- Activité une\n- Activité deux\n- Activité trois\n- Activité quatre\n"
            "Compétences : SQL, Power BI\nCOMPETENCES CLES\n- Analyse statistique\n"
            "APTITUDES\n- Esprit d'équipe\n"
        )
        source = candidatures_app.make_cv_latex(profile, application, draft)
        self.assertIn(r"\item Activité une", source)
        self.assertIn(r"\item Activité deux", source)
        self.assertNotIn(r"\item Activité trois", source)
        self.assertNotIn(r"\item Activité quatre", source)
        self.assertNotIn(r"\cvsection{Compétences clés}", source)
        self.assertNotIn(r"\cvsection{Aptitudes}", source)

    def test_cv_experiences_are_sorted_newest_to_oldest_across_formats(self):
        profile = {
            "full_name": "Nom Exemple",
            "email": "",
            "phone": "",
            "location": "",
        }
        application = {"position": "Data Analyst", "company": "Exemple"}
        draft = (
            "PROFIL\nAnalyste de données.\n"
            "EXPERIENCES PROFESSIONNELLES\n"
            "Data Scientist 03/2025 – 09/2025\nPOKAPOK\n"
            "- Mission POKAPOK.\n"
            "Data Analyst, Génétique Statistique\nUniversité d'Angers\n10/2025 – 02/2026\n"
            "- Mission Université.\n"
            "Data Analyst & BI Specialist 04/2026 – 09/2026\nParker Aerospace\n"
            "- Mission Parker.\n"
            "FORMATION\nMaster Data\n"
        )

        sorted_draft = candidatures_app.sort_cv_experience_sections(draft)
        self.assertLess(
            sorted_draft.index("Parker Aerospace"),
            sorted_draft.index("Université d'Angers"),
        )
        self.assertLess(
            sorted_draft.index("Université d'Angers"),
            sorted_draft.index("POKAPOK"),
        )

        source = candidatures_app.make_cv_latex(profile, application, draft)
        self.assertLess(
            source.index(r"\cventry{Data Analyst \& BI Specialist}"),
            source.index(r"\cventry{Data Analyst, Génétique Statistique}"),
        )
        self.assertLess(
            source.index(r"\cventry{Data Analyst, Génétique Statistique}"),
            source.index(r"\cventry{Data Scientist}"),
        )

        docx = Document(io.BytesIO(
            candidatures_app.make_cv_docx(profile, application, draft).getvalue()
        ))
        docx_text = "\n".join(paragraph.text for paragraph in docx.paragraphs)
        self.assertLess(docx_text.index("Parker Aerospace"), docx_text.index("Université d'Angers"))
        self.assertLess(docx_text.index("Université d'Angers"), docx_text.index("POKAPOK"))

        pdf = PdfReader(io.BytesIO(
            candidatures_app.make_pdf(profile, application, draft, "cv").getvalue()
        ))
        pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        self.assertLess(pdf_text.index("Parker Aerospace"), pdf_text.index("Université d'Angers"))
        self.assertLess(pdf_text.index("Université d'Angers"), pdf_text.index("POKAPOK"))

    def test_cv_latex_splits_semicolon_education_and_fits_one_page(self):
        profile = {
            "full_name": "Moubarak LASSISSI",
            "location": "Angers",
            "email": "moubarak@example.com",
            "phone": "+33000000000",
            "certifications": (
                "AWS Certified Machine Learning - Specialty\nDocker and Kubernetes\n"
                "Foundations of Deep Reinforcement Learning with PyTorch\n"
                "Generative AI and LLMs: Architecture and Data Preparation\nWorking with SSH\n"
                "Databases and SQL for Data Science with Python\nPython with Jupyter Notebooks\n"
                "Power BI Bootcamp - Learn by Building a Real-World Project\n"
                "Azure DevOps and Continuous Delivery with Git\nAzure Devops Boards Projects\n"
                "Alteryx Mastery: Data Analytics & Workflow Automation\n"
                "Data Modeling with Power BI: Advanced DAX Calculations\n"
                "Creative Designing in Power BI\n"
                "Identifiant du diplôme : Creative Designing in Power BI\n"
                "Python for Data Science, AI & Development\n"
                "Developing AI Applications with Python and Flask"
            ),
            "languages": "Français : Langue maternelle\nAnglais : Professionnel",
        }
        application = {"position": "Data Scientist", "company": "Baker Tilly"}
        draft = """PROFIL
Data Scientist spécialisé en modélisation statistique et analyse de données. Expert R et SQL, maîtrise du traitement de données volumétriques et exploitation de sources administratives.
COMPETENCES CLES
- Conception et déploiement de modèles Machine Learning
- Analyse statistique avancée
- Préparation et structuration de données
- Conception de tableaux de bord interactifs
- Développement avec R, Python, SQL et Power BI
- Industrialisation avec Git, MLflow et Docker
EXPERIENCES PROFESSIONNELLES
DATA ANALYST & BI SPECIALIST 04/2026 – 09/2026
Parker Aerospace
- Produire et assurer la fiabilité des indicateurs de gestion et du reporting.
- Coordonner avec les équipes métier pour définir les besoins.
- Concevoir des tableaux de bord interactifs avec Power BI.
DATA ANALYST, GÉNÉTIQUE STATISTIQUE 10/2025 – 02/2026
Université d'Angers
- Analyser des données génétiques volumineuses.
- Réaliser des tests statistiques avancés.
- Développer du code statistique avec R.
DATA SCIENTIST 03/2025 – 09/2025
POKAPOK & IMT Atlantique
- Réaliser des analyses statistiques et modélisations prédictives.
- Industrialiser les analyses avec Git, MLflow et Docker.
- Analyser les erreurs de prédiction des modèles.
STATISTICIEN | SUIVI BUDGÉTAIRE 01/2021 – 12/2022
Centre d'Enseignement Général
- Gérer le suivi budgétaire et l'analyse des dépenses.
- Construire des indicateurs et tableaux de bord.
FORMATION
MASTER EN DATA SCIENCES 2025 – 2026 ; MASTER EN INGÉNIERIE STATISTIQUE ; Nantes Universités ; 2023 – 2025 ; LICENCE EN MATHÉMATIQUES ; Université de Brest ; 2022 – 2023 ; LICENCE MATHÉMATIQUES STATISTIQUES ; Université de Kara ; 2017 – 2020 ; APTITUDES
- Communication claire
- Flexibilité et adaptation
- Esprit d'équipe
LOGICIELS
- R
- Python
- SQL
- Power BI
- Git
- Spark
- Azure
"""
        source = candidatures_app.make_cv_latex(profile, application, draft)
        self.assertIn(r"\cventry{LICENCE EN MATHÉMATIQUES}{Université de Brest}{2022 – 2023}", source)
        self.assertNotIn(r"\cvsection{Aptitudes}", source)
        self.assertNotIn(r"\cvsection{Compétences clés}", source)
        self.assertLess(source.index(r"\cvsection{Profil}"), source.index(r"\cvsection{Expériences professionnelles}"))
        self.assertNotIn("AWS Certified Machine Learning", source)
        self.assertNotIn("Developing AI Applications with Python and Flask", source)
        self.assertNotIn("Identifiant du diplôme", source)
        if candidatures_app.shutil.which("xelatex"):
            pdf = PdfReader(candidatures_app.compile_latex(source))
            self.assertEqual(len(pdf.pages), 1)
            text = pdf.pages[0].extract_text() or ""
            self.assertIn("Université de Kara", text)
            self.assertNotIn("Communication claire", text)

    def test_cv_includes_only_certifications_relevant_to_job(self):
        profile = {
            "full_name": "Nom", "email": "", "phone": "", "location": "",
            "certifications": (
                "AWS Certified Machine Learning - Specialty\nDocker and Kubernetes\n"
                "Foundations of Deep Reinforcement Learning with PyTorch\n"
                "Generative AI and LLMs: Architecture and Data Preparation\n"
                "Working with SSH\nDatabases and SQL for Data Science with Python\n"
                "Power BI Bootcamp - Learn by Building a Real-World Project\n"
                "Alteryx Mastery: Data Analytics & Workflow Automation"
            ),
        }
        application = {
            "position": "Data Analyst",
            "job_description": "Créer des tableaux de bord Power BI et préparer les données avec SQL.",
        }
        source = candidatures_app.make_cv_latex(
            profile, application, "PROFIL\nAnalyste de données.\n"
        )
        self.assertIn("Power BI Bootcamp", source)
        self.assertIn("Databases and SQL for Data Science", source)
        self.assertNotIn("AWS Certified", source)
        self.assertNotIn("Working with SSH", source)
        self.assertEqual(
            candidatures_app.select_relevant_items(
                candidatures_app.profile_list(profile, "certifications"), application
            ),
            [
                "Power BI Bootcamp - Learn by Building a Real-World Project",
                "Databases and SQL for Data Science with Python",
            ],
        )

    def test_cv_project_public_link_is_clickable_in_exports(self):
        profile = {
            "full_name": "Nom Exemple",
            "email": "",
            "phone": "",
            "location": "",
        }
        application = {"position": "Data Analyst", "company": "Exemple"}
        url = "https://github.com/exemple/solarwatch"
        content = (
            "PROJETS\nSolarWatch\nProjet personnel\n"
            "- Résultat documenté : alertes candidates.\n"
            f"- À consulter : {url}"
        )

        latex = candidatures_app.make_cv_latex(profile, application, content)
        self.assertIn(rf"\href{{\detokenize{{{url}}}}}", latex)
        if candidatures_app.shutil.which("xelatex"):
            pdf_output = candidatures_app.compile_latex(latex, one_page=True)
            pdf = PdfReader(pdf_output)
            self.assertEqual(len(pdf.pages), 1)
            self.assertTrue(any(
                annotation.get_object().get("/A", {}).get("/URI") == url
                for annotation in pdf.pages[0].get("/Annots", [])
            ))

        docx_output = candidatures_app.make_cv_docx(profile, application, content)
        with zipfile.ZipFile(docx_output) as archive:
            relationships = archive.read("word/_rels/document.xml.rels").decode()
        self.assertIn(url, relationships)
        self.assertIn("TargetMode=\"External\"", relationships)

        pdf = PdfReader(io.BytesIO(candidatures_app.make_pdf(profile, application, content, "cv").getvalue()))
        annotations = pdf.pages[0].get("/Annots", [])
        self.assertTrue(any(
            annotation.get_object().get("/A", {}).get("/URI") == url
            for annotation in annotations
        ))

    def test_cv_keeps_only_tools_matching_the_job(self):
        profile = {
            "full_name": "Nom", "email": "", "phone": "", "location": "",
        }
        application = {
            "position": "Data Analyst",
            "job_description": "Créer des tableaux de bord Power BI avec DAX et SQL.",
        }
        draft = (
            "PROFIL\nAnalyste de données.\nLOGICIELS\n"
            "Power BI · SQL · Python · R · Excel · Power Query · DAX · Git · Docker · Azure DevOps"
        )
        source = candidatures_app.make_cv_latex(profile, application, draft)
        self.assertIn("Power BI", source)
        self.assertIn("SQL", source)
        self.assertIn("DAX", source)
        self.assertNotIn("Docker", source)
        self.assertNotIn("Azure DevOps", source)

    def test_cv_long_content_is_scaled_to_one_page(self):
        if not candidatures_app.shutil.which("xelatex"):
            self.skipTest("XeLaTeX is not installed")
        profile = {
            "full_name": "Nom Exemple", "email": "nom@example.com",
            "phone": "", "location": "Angers",
        }
        application = {"position": "Data Analyst", "company": "Exemple"}
        content = "PROFIL\n" + "\n".join(
            f"Expérience détaillée numéro {index} avec des missions d'analyse, "
            "de reporting et de traitement de données décrites avec précision."
            for index in range(1, 45)
        )
        pdf = PdfReader(io.BytesIO(
            candidatures_app.compile_latex(
                candidatures_app.make_cv_latex(profile, application, content)
            ).getvalue()
        ))
        self.assertEqual(len(pdf.pages), 1)

    def test_cv_fallback_pdf_keeps_long_content_to_one_page(self):
        profile = {
            "full_name": "Nom Exemple", "email": "nom@example.com",
            "phone": "", "location": "Angers",
        }
        application = {"position": "Data Analyst", "company": "Exemple"}
        content = "PROFIL\n" + "\n".join(
            f"Mission détaillée numéro {index} avec analyse, reporting et données."
            for index in range(1, 45)
        )
        pdf = PdfReader(io.BytesIO(
            candidatures_app.make_pdf(profile, application, content, "cv").getvalue()
        ))
        self.assertEqual(len(pdf.pages), 1)


    def test_profile_sections_and_experience_skills_in_latex(self):
        profile = {
            "full_name": "A", "email": "", "phone": "", "location": "",
            "languages": "Français – natif\nAnglais – B2", "hobbies": "Course; Échecs",
            "certifications": "PL-300 – 2025", "achievements": "",
        }
        draft = (
            "PROFIL\nTexte\nEXPERIENCES PROFESSIONNELLES\nDATA ANALYST 04/2026 – 09/2026\nACME\n"
            "- Réduit le reporting de 40 %\nCompétences : SQL, Power BI\n"
            "LANGUES\nChinois\n"
        )
        source = candidatures_app.make_cv_latex(
            profile, {"position": "Data Analyst", "company": "X"}, draft
        )
        self.assertIn(r"\cvskills{SQL, Power BI}", source)
        self.assertIn("Anglais – B2", source)
        self.assertNotIn("Chinois", source)
        self.assertIn(r"\cvsection{Certifications}", source)
        self.assertIn(r"\cvsection{Loisirs}", source)
        self.assertIn("Échecs", source)
        if candidatures_app.shutil.which("xelatex"):
            self.assertEqual(len(PdfReader(candidatures_app.compile_latex(source)).pages), 1)

    def test_cv_layout_is_validated_and_applied(self):
        self.assertEqual(candidatures_app.parse_cv_layout('{"font": 99, "accent": "zzz"}')["font"], 8.5)
        profile = {"full_name": "A", "email": "", "phone": "", "location": ""}
        application = {"position": "Data", "company": "X"}
        source = candidatures_app.make_cv_latex(
            profile, application, "PROFIL\nTexte\n", {"font": 9.5, "margin": 2, "accent": "1F3A5F"}
        )
        self.assertIn("margin=2.0cm", source)
        self.assertIn("{HTML}{1F3A5F}", source)
        self.assertNotIn("@@", source)


if __name__ == "__main__":
    unittest.main()
